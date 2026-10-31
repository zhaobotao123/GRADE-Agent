from pathlib import Path
from copy import deepcopy
import json,re,csv,io,zipfile,hashlib
from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement
from lxml import etree

BASE=Path(__file__).resolve().parents[1]
F=Path('D:/Users/84130/Desktop/投稿/转投修改_20261022')
QA=BASE/'outputs/submission_audit_v33_20261028'
d=Document(F/'GRADE-Agent_SEE_Restructured_v32.docx')
math=[etree.tostring(x)for x in d._element.xpath('.//m:oMath')]
tables=[t._tbl.xml for t in d.tables]
changes=[]
def replace(p,text):
    before=p.text
    assert not p._p.xpath('.//m:oMath|.//w:drawing')
    prop=deepcopy(p.runs[0]._r.rPr)if p.runs else None
    p.clear();r=p.add_run(text)
    if prop is not None:r._r.insert(0,prop)
    changes.append({'before':before,'after':text})
def bystart(start):return next(p for p in d.paragraphs if p.text.startswith(start))
for p in d.paragraphs:
    if 'ordinary multi-turn LLM' in p.text:replace(p,p.text.replace('ordinary multi-turn LLM','ordinary LLM'))
replace(bystart('4.2 '),'4.2 Predictive Contributions of Student and Contextual Characteristics')
replace(bystart('Mathematics achievement was represented'),"Mathematics achievement was represented by ten plausible values (PV1MATH–PV10MATH). We fitted the final prediction models and DML analyses separately for each plausible value. Their mean (MATH_MEAN) supported descriptive reporting and training-set hyperparameter selection. It was not the outcome for the final plausible-value-specific analyses.")
replace(bystart('We specified 12 candidate factors'),"The analysis configuration specified 14 candidate factors before SHAP ranking, based on educational relevance and available PISA constructs. FEELSAFE and REPEAT exceeded the 50% missingness threshold, leaving 12 factors for DML. Thus, data availability reduced the initial list; SHAP rankings did not determine the retained factors.")
replace(bystart('DML adjustment included'),"For each focal factor, DML adjusted for the other 11 retained factors and nine background variables: CNT, ST004D01T, AGE, GRADE, LANGN, HISCED, PAREDINT, HISEI, and RELATST. CNT was constant within Japan. This 21-variable DML pool was separate from the expanded 42-variable prediction set; identifiers and weights were excluded as predictors.")
replace(bystart('We imputed missing numerical'),"We imputed missing numerical values using training-data medians and categorical values using the most frequent category, followed by one-hot encoding. Linear prediction models used standardised transformed inputs. DML excluded observations missing the focal factor, outcome, positive weight, or school identifier, then fitted covariate preprocessing within each cross-fitting fold.")
replace(bystart('We compared Random Forest'),"We compared Random Forest, ExtraTrees, RidgeCV, ElasticNetCV, and LightGBM (Ke et al., 2017) using the same 42 original predictors and 80%/20% split. Three-fold cross-validation within the training partition selected tree-model settings using MATH_MEAN and weighted predictive performance. ElasticNetCV additionally used internal five-fold tuning; RidgeCV used its default internal selection. Final evaluation refitted each selected specification separately for all ten plausible values.")
replace(bystart('We calculated SHAP values'),"We calculated SHAP values for the selected LightGBM model on up to 3,000 uniformly sampled training observations per plausible value. We weighted absolute SHAP values using student weights, summed encoded contributions back to original variables, and averaged the results across ten analyses.")
replace(bystart('Figure 1. '),'Figure 1. Study framework linking achievement-related analysis, conversational educational support, and teacher evaluation. Population-level predictive contributions and adjusted associations inform discussion of user-supplied learning circumstances. Two teachers evaluated 120 English final replies across 12 synthetic cases and five runs per system.')
replace(bystart('Figure 2. '),'Figure 2. Predictive contributions and adjusted associations informing GRADE-Agent. (a) Weighted performance of five prediction models. (b) Student and contextual characteristics ranked by mean absolute SHAP value. (c) Pooled DML adjusted associations with 95% normal-approximation confidence intervals, using school-clustered standard errors and plausible-value pooling. (d) Organisation of findings for contextual interpretation. Predictive contributions and adjusted associations do not establish individual causes.')
replace(bystart('Figure 4. '),'Figure 4. Matched comparison and blinded teacher evaluation of GRADE-Agent and an ordinary LLM.')
replace(bystart('Table 2. '),'Table 2. Teacher ratings of final replies from GRADE-Agent and an ordinary LLM')
# Restore explicit method navigation while preserving existing editable equations.
for prefix,title in [('We compared Random Forest','3.3.1 Weighted Predictive Modelling'),('We calculated SHAP values','3.3.2 SHAP Feature Attribution'),('We applied DML separately','3.3.3 Double Machine Learning')]:
    anchor=bystart(prefix);el=OxmlElement('w:p');anchor._p.addprevious(el);p=Paragraph(el,anchor._parent);p.style=d.styles['Heading 3'];p.add_run(title)
# Replace an over-specific conversational claim with what the cited reviews support.
p=bystart('Conversation allows users')
replace(p,"Reviews describe conversational interaction and personalised support in educational chatbots and discuss opportunities and limitations of educational AI (Chiu et al., 2023; Kuhail et al., 2023; Yan et al., 2024). Here, these ideas motivate contextual inquiry before a summary. Educators judge whether interpretations fit the learner and which next steps are appropriate.")
assert math==[etree.tostring(x)for x in d._element.xpath('.//m:oMath')]
assert tables==[t._tbl.xml for t in d.tables]
target=F/'GRADE-Agent_SEE_Audited_Draft_v33.docx';d.save(target)
# Synchronise only the current statistical analysis, not frozen generation seeds.
original=F/'Supplementary_Evaluation_Materials_v21.zip'
with zipfile.ZipFile(original)as z:items={n:z.read(n)for n in z.namelist()}
result=json.loads((BASE/'outputs/teacher_bootstrap_seed42_20261028/results_seed42.json').read_text(encoding='utf-8'))
items['results.json']=json.dumps(result,ensure_ascii=False,indent=2).encode()
buffer=io.StringIO();writer=csv.DictWriter(buffer,fieldnames=list(result['summary'][0]));writer.writeheader();writer.writerows(result['summary'])
items['results_summary.csv']=buffer.getvalue().encode('utf-8-sig')
items['analysis/reproduce_statistics.py']=items['analysis/reproduce_statistics.py'].replace(b'default_rng(20261027)',b'default_rng(42)')
oldname='analysis/finalize_teacher_results_v21.py'
oldcode=items.pop(oldname)
items['analysis/finalize_teacher_results_seed42.py']=oldcode.replace(b'default_rng(20261027)',b'default_rng(42)').replace(b"'seed':20261027",b"'seed':42")
readme=items['README.md'].decode().replace('seed 20261027','seed 42').replace('finalize_teacher_results_v21.py','finalize_teacher_results_seed42.py')
readme+='\n\nAnalysis revision: bootstrap seed changed to 42 at the author\'s request. The previous analysis is retained separately. Frozen generation and presentation-order seeds remain unchanged. No original replies or teacher scores were modified.\n'
items['README.md']=readme.encode()
manifest={n:hashlib.sha256(v).hexdigest()for n,v in items.items()if n!='SHA256SUMS.json'}
items['SHA256SUMS.json']=json.dumps(manifest,indent=2).encode()
archive=F/'Supplementary_Evaluation_Materials_v33_seed42.zip'
with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED)as z:
    for n,v in items.items():z.writestr(n,v)
ri=next(i for i,p in enumerate(d.paragraphs)if p.text=='References')
intro=next(i for i,p in enumerate(d.paragraphs)if p.text=='1. Introduction')
conclusionend=next(i for i,p in enumerate(d.paragraphs)if p.text=='Data Availability')
count=lambda text:len(re.findall(r"\b[\w]+(?:[’'–-][\w]+)*\b",text))
counts={'main_text_including_headings_and_captions':count('\n'.join(p.text for p in d.paragraphs[intro:conclusionend])), 'full_document_including_references_and_tables':count('\n'.join(p.text for p in d.paragraphs)+'\n'+'\n'.join(c.text for t in d.tables for row in t.rows for c in row.cells))}
(QA/'v33_revision.json').write_text(json.dumps({'target':str(target),'changes':changes,'counts':counts,'equations_preserved':True,'tables_preserved':True,'figure_internal_labels_still_pending':True},ensure_ascii=False,indent=2),encoding='utf-8')
report=['# 投稿核查：已完成与待完成', '', '## 已核实的统计', '',
'五种模型×十个合理值的50行指标已重算均值和标准差；42个原始预测变量×十个合理值的420行SHAP记录已核对；12个因素×十个合理值的120行DML估计已重算合并系数和标准误。两教师、两系统、12案例、5次运行构成120条回复及960个评分，均值、配对差异、种子42的50000次案例Bootstrap区间及加权κ独立复算一致。', '',
'## 必须修正的方法描述（已写入v33草稿）', '',
'- 树模型调参为训练集内三折，不是五折；ElasticNetCV另有内部五折。',
'- MATH_MEAN用于描述和调参，不仅用于描述；最终预测与DML仍逐合理值估计。',
'- 初始14个DML候选经缺失率筛选留下12个；并非最初只列12个。',
'- DML实际使用21变量池，每个焦点因素对应其余20个调整变量；不是扩展的42变量预测集。',
'- 类别缺失值填补实际为众数；焦点因素缺失的DML观测直接排除。',
'- SHAP是均匀抽样后按学生权重汇总，不是按权重抽样。', '',
'## 已完成的文字修改', '',
'统一正文ordinary LLM；更新4.2与图1、图2、图4、表2标题；3.3恢复三种方法的小节导航，全部原生DML公式保留。新补充包已同步统计随机种子42；生成与匿名呈现的历史种子不改。原始评分、回复、旧版本不覆盖。', '',
'## 词数（正则计数，Word可能略有差异）', '',str(counts),'',
'## 参考文献核查范围', '',
'26条参考文献中23条提供DOI，均取得CrossRef元数据；这不等于所有作者细节和论述支持均已全文验证。Ke和Lundberg需按NeurIPS会议原始记录核对；Hu in press仍依作者提供资料，未独立确认出版信息。OECD2010、2023a、2023b及2024的DOI均指向对应出版物；2023a为Volume I，2023b为Volume II。', '',
'原始元数据及统计复算记录已保存。不能把本轮标成“全部参考文献论述均已核验”。', '',
'## 尚未完成，不能宣称可直接投稿', '',
'1. 图内ordinary multi-turn LLM等文字尚未从绘图源同步；正文图注已改。',
'2. 逐条全文支持核查和完整引文双向匹配尚未结束。',
'3. 官网作者指南返回403，最新字数要求未能由官方页面确认。历史材料中9000字和150字摘要限制不能冒充当前官方规则。',
'4. 直接读取大型PISA原始文件遇到内存错误，未取得新的原始样本核查输出；5760由原分析日志和描述表支持。',
'5. 摘要最终精简、投稿Highlights和补充章节编号的完整核对仍待完成。',
'6. 图3旧报告界面与当前段落总结输出的对应仍需确认。',
'7. 作者要求暂缓作者信息及贡献。经费尚未确定；不得填无资助。利益冲突未确认；不得填无冲突。教师仅确认匿名评分，伦理审批/豁免及知情同意尚未确认，不代写已审批。', '',
'官方数据源：https://www.oecd.org/en/data/datasets/pisa-2022-database.html',
'官方作者指南：https://www.sciencedirect.com/journal/studies-in-educational-evaluation/publish/guide-for-authors',
'Highlights一般格式参考：https://www.elsevier.support/publishing/answer/how-do-i-include-highlights-with-my-manuscript']
(F/'GRADE-Agent_投稿核查进度_v33.md').write_text('\n'.join(report),encoding='utf-8')
print(target);print(archive);print(counts)

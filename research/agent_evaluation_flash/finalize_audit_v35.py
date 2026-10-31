"""Complete the source audit without changing experimental records."""
import json,re,zipfile,hashlib,subprocess,sys
from pathlib import Path
from copy import deepcopy
from docx import Document
from lxml import etree
BASE=Path(__file__).resolve().parents[1]
F=Path('D:/Users/84130/Desktop/投稿/转投修改_20261022')
OUT=BASE/'outputs/submission_audit_v35_20261028';OUT.mkdir(parents=True,exist_ok=True)
d=Document(F/'GRADE-Agent_SEE_Checked_v34.docx')
math=[etree.tostring(x)for x in d._element.xpath('.//m:oMath')]
tables=[t._tbl.xml for t in d.tables]
changes=[]
def replace(p,t):
 old=p.text; prop=deepcopy(p.runs[0]._r.rPr)if p.runs else None
 p.clear();r=p.add_run(t)
 if prop is not None:r._r.insert(0,prop)
 changes.append({'before':old,'after':t})
ri=next(i for i,p in enumerate(d.paragraphs)if p.text=='References')
for i,p in enumerate(d.paragraphs):
 t=p.text
 if i<ri:
  t=t.replace('SHAP feature attribution, and Double Machine Learning to examine','feature attribution, and Double Machine Learning (DML) to examine')
  t=t.replace('an ordinary large language model across twelve','an ordinary large language model (LLM) across twelve')
  t=t.replace('fell outside the 12 prespecified DML factors','fell outside the 12 retained DML factors')
  if t.startswith('Here,') and 'three-fold cross-fitting' in t:
   t='Here, Yᵢ denotes a mathematics plausible value, Dᵢ is the candidate factor under examination, and Xᵢ contains the prespecified adjustment variables. We estimated the conditional means ℓ(Xᵢ) = E(Yᵢ | Xᵢ) and m(Xᵢ) = E(Dᵢ | Xᵢ) using LightGBM with three-fold cross-fitting. The resulting weighted residual-on-residual estimator was:'
 else:
  if t.startswith('Albreiki,'):t=t.replace("students' performance prediction",'student’ performance prediction')
  if t.startswith('Alnasyan,'):t=t.replace('Alassafi, M. O.','Alassafi, M.')
  if t.startswith('Black,'):t=t.replace('Accountability, 21,','Accountability, 21(1),')
  if t.startswith('Bond,'):t=t.replace('Higher Education, 21,','Higher Education, 21(1),')
  if t.startswith('Chernozhukov,'):t=t.replace('Journal, 21,','Journal, 21(1),')
  if t.startswith('Khosravi,'):t=t.replace('Shum, S. B.','Buckingham Shum, S.').replace('Tsai, Y. S.','Tsai, Y.-S.')
  if t.startswith('Kuhail,'):t=t.replace('Technologies, 28,','Technologies, 28(1),')
  if t.startswith('Lundberg,'):
   t=t.replace('Lee, S. I.','Lee, S.-I.')
   t+=' https://proceedings.neurips.cc/paper_files/paper/2017/hash/8a20a8621978632d76c43dfd28b67767-Abstract.html'
  if t.startswith('Shmueli,'):t=t.replace('Science, 25,','Science, 25(3),')
  if t.startswith('Yan,'):t=t.replace('Behaviour, 8,','Behaviour, 8(10),')
  if t.startswith('Zawacki-Richter,'):t=t.replace('Higher Education, 16,','Higher Education, 16(1),')
  t=t.replace('Computers & Education: Artificial Intelligence','Computers and Education: Artificial Intelligence')
 if t!=p.text:replace(p,t)
assert tables==[t._tbl.xml for t in d.tables]
# Distinguish the structural g(X) in equation1 from the outcome conditional mean
# used for residualisation in equation3. Native OMML remains editable.
equations=d._element.xpath('.//m:oMath')
math=[etree.tostring(e)for e in equations]
changed_math=[]
for j,e in enumerate(equations):
 texts=[t for t in e.iter()if etree.QName(t).localname=='t']
 if any(etree.QName(t).localname=='nary'for t in e.iter()) and any(t.text=='g'for t in texts):
  for t in texts:
   if t.text=='g':t.text='ℓ'
  changed_math.append(j)
assert len(changed_math)==1,changed_math
assert all(old==etree.tostring(e)for j,(old,e)in enumerate(zip(math,equations))if j not in changed_math)
# Keep journal/proceedings names and volume numbers consistently italic after
# metadata replacements, which otherwise flatten mixed-format Word runs.
journals=['Education Sciences','Educational Assessment, Evaluation and Accountability','International Journal of Educational Technology in Higher Education','The Econometrics Journal','Computers and Education: Artificial Intelligence','Transactions of the Association for Computational Linguistics','Studies in Educational Evaluation','Teaching and Teacher Education','Education and Information Technologies','Statistical Science','Nature Human Behaviour','Advances in Neural Information Processing Systems']
for p in d.paragraphs[ri+1:]:
 text=p.text; spans=[]
 for journal in journals:
  match=re.search(re.escape(journal)+r'(?:, (?:\d+|[A-Z]\d+))?',text)
  if match:spans.append(match.span())
 if text.startswith('OECD.'):
  begin=text.index('). ')+3;end=text.index('. OECD Publishing')
  spans.append((begin,end))
 if spans:
  prop=deepcopy(p.runs[0]._r.rPr)if p.runs else None
  p.clear();pos=0
  for begin,end in sorted(spans):
   for content,it in [(text[pos:begin],False),(text[begin:end],True)]:
    if content:
     r=p.add_run(content)
     if prop is not None:r._r.insert(0,deepcopy(prop))
     r.italic=it
   pos=end
  if pos<len(text):
   r=p.add_run(text[pos:])
   if prop is not None:r._r.insert(0,deepcopy(prop))
   r.italic=False
target=F/'GRADE-Agent_SEE_Checked_v35.docx';d.save(target)
pars=d.paragraphs;ri=next(i for i,p in enumerate(pars)if p.text=='References')
body='\n'.join(p.text for p in pars[:ri]);refs=[p.text for p in pars[ri+1:]if p.text.strip()]
words=lambda s:len(re.findall(r"\b[\w]+(?:[’'-][\w]+)*\b",s))
intro=next(i for i,p in enumerate(pars)if p.text=='1. Introduction')
abstract=pars[next(i for i,p in enumerate(pars)if p.text=='Abstract')+1].text
counts={'abstract_words_space_delimited':len(abstract.split()),'main_words_including_headings_and_captions_excluding_abstract_references_tables':words('\n'.join(p.text for p in pars[intro:ri])),
 'all_words_including_abstract_references_tables':words('\n'.join(p.text for p in pars)+'\n'+'\n'.join(c.text for t in d.tables for row in t.rows for c in row.cells))}
metadata=json.loads((BASE/'outputs/submission_audit_v33_20261028/reference_metadata.json').read_text(encoding='utf-8'))['references']
support={
'Albreiki':('作者摘要/元数据','学生表现机器学习预测的一般背景成立；不外推为本系统或日本学生的验证。','https://api.crossref.org/works/10.3390/educsci11090552'),
'Alnasyan':('出版者摘要','46项虚拟学习环境研究支持预测应用背景，不作为个体因果判断依据。','https://www.sciencedirect.com/science/article/pii/S2666920X24000328'),
'Black':('作者机构全文','形成性评价定义明确涉及收集、解释和使用信息决定教学下一步；作者稿PDF第8页。','https://kclpure.kcl.ac.uk/ws/portalfiles/portal/9119063/Black2009_Developing_the_theory_of_formative_assessment.pdf'),
'Bond':('出版者摘要','支持伦理、合作和方法严谨性的研究需求；未声称本系统已满足全部要求。','https://doi.org/10.1186/s41239-023-00436-z'),
'Chernozhukov':('作者预印本摘要/公式依据','确认正交化、辅助预测与交叉拟合；本研究实现以本地代码为准，因果识别假设并未由算法自动保证。','https://arxiv.org/abs/1608.00060'),
'Chiu':('出版者摘要','支持教育AI的应用、机会与挑战概述，未用于证明本系统实际效果。','https://www.sciencedirect.com/science/article/pii/S2666920X2200073X'),
'Fabbri':('官方全文','Section4.3/PDF第6页定义一致性、相关性、流畅性与连贯性；本研究是适配量表，非原量表原封验证。','https://aclanthology.org/2021.tacl-1.24.pdf'),
'Hu':('作者提供，外部未确认','研究内容依据已提供匿名手稿；in press和完整作者顺序仍依据作者确认，不补造DOI。','作者提供的接收/书目信息'),
'Jin':('作者机构摘要','支持在线同伴评价中的行为模式与表现预测描述。','https://repository.eduhk.hk/en/publications/investigating-behavioral-patterns-to-facilitate-performance-predi/'),
'Ke':('原始会议记录','确认LightGBM方法、年份、全部8位作者和会议卷号。','https://proceedings.neurips.cc/paper_files/paper/2017/hash/6449f44a102fde848669bdd9eb6b76fa-Abstract.html'),
'Khosravi':('出版者全文','2.3、5.3、5.4支持教育情境、利益相关者和解释评价；第二作者复姓按本人出版记录为Buckingham Shum。','https://doi.org/10.1016/j.caeai.2022.100074'),
'Klug':('出版者摘要及节选','教师诊断过程模型及N=293案例验证，支持观察、解释和教学决策的概念联系。','https://www.sciencedirect.com/science/article/abs/pii/S0742051X12001503'),
'Kuhail':('出版者摘要','36项研究支持教育聊天机器人互动设计背景，不证明追问在本试验中有效。','https://doi.org/10.1007/s10639-022-11177-3'),
'Lee':('出版者摘要及方法节选','支持大型在线课程的AI辅助反馈评价描述。','https://doi.org/10.1016/j.stueduc.2023.101250'),
'Liu':('官方摘要','G-Eval是GPT4自动评估；本文明确没有实施自动G-Eval，本条仅支持文本评价维度背景。','https://aclanthology.org/2023.emnlp-main.153/'),
'Lundberg':('原始会议记录','确认SHAP与输入贡献解释；修正Su-In Lee缩写并补原会议链接。','https://proceedings.neurips.cc/paper_files/paper/2017/hash/8a20a8621978632d76c43dfd28b67767-Abstract.html'),
'Martínez-Abad':('出版者摘要及方法节选','西班牙PISA2015学校效能研究支持引言表述。','https://www.sciencedirect.com/science/article/pii/S0191491X20301231'),
'Shmueli':('作者摘要','支持预测与解释目标区分。','https://doi.org/10.1214/10-STS330'),
'Yan':('作者机构摘要','支持生成式AI学习应用的机会、挑战和跨学科视角；未声称其验证本系统。','https://orbilu.uni.lu/handle/10993/62522'),
'Zawacki-Richter':('出版者摘要','支持教育者参与、教学与伦理视角的研究缺口。','https://doi.org/10.1186/s41239-019-0171-0'),
}
rows=[]
for r in metadata:
 a,y=r['first_author'],r['year'];hit=bool(re.search(re.escape(a)+r'[^;\n]{0,70}?'+re.escape(y),body))
 if a=='OECD'and y=='2023b':hit=hit or 'OECD, 2023a, 2023b'in body
 assert hit,(a,y)
 if a=='OECD':
  desc='出版者书目与数据记录；2023a=VolumeI，2023b=VolumeII；2024=技术报告。'
  scope='官方记录'
  if y=='2010':scope='OECD官方全文';desc='结论部分明确提出更多支持可能投向低成绩学生；只支持可能解释，不证明本研究反向因果。'
  entry=(scope,desc,'https://doi.org/'+r['doi'])
 elif a=='Wang':
  entry=('作者机构摘要'if y=='2026'else'出版者全文节选','57项研究及理论/方法需求已核实。'if y=='2026'else'方法、结果和讨论区分ICT资源与学生参与/实际使用，支持6.2限定性表述。','https://research.cuhk.edu.hk/en/publications/the-application-of-machine-learning-approaches-in-international-l/'if y=='2026'else'https://www.sciencedirect.com/science/article/pii/S0191491X23000743')
 else:entry=support[a]
 rows.append({'reference':a+' '+y,'in_text':True,'scope':entry[0],'assessment':entry[1],'source':entry[2]})
# Citation occurrences were inspected paragraph by paragraph; check every observed author-year pair.
known={(r['first_author'],r['year'])for r in metadata}
observed=[('Martínez-Abad','2020'),('Wang','2026'),('Wang','2023'),('Jin','2024'),('Shmueli','2010'),('Lundberg','2017'),('Chernozhukov','2018'),('Lee','2023'),('Hu','in press'),('OECD','2024'),('Albreiki','2021'),('Alnasyan','2024'),('Black','2009'),('Klug','2013'),('Bond','2024'),('Khosravi','2022'),('Zawacki-Richter','2019'),('Chiu','2023'),('Kuhail','2023'),('Yan','2024'),('OECD','2023a'),('OECD','2023b'),('Ke','2017'),('Fabbri','2021'),('Liu','2023'),('OECD','2010')]
assert set(observed)==known
assert len(refs)==26
assert not re.search(r'ordinary multi-turn LLM|adjusted-effect estimation|adjusted-effect evidence',body,re.I)
assert len(re.findall(r'^Figure [1-5]\.',body,re.M))==5
audit={'changes':changes,'counts':counts,'reference_support':rows,'citations_in_both_directions_match':True,'tables_unchanged':True,'native_formula3_nuisance_symbol_corrected':changed_math,'raw_sample_audit':json.loads((BASE/'outputs/submission_audit_v33_20261028/raw_sample_audit.json').read_text(encoding='utf-8'))}
(OUT/'final_audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf-8')
report=['# 论文核查与修订记录 v35','','作者、贡献和经费信息按要求暂不填写；其余核查独立进行。保留全部旧版本和原始评分，未为本轮核查调用大模型或改动实验回复。','','## 1 结构与术语','','- ordinary LLM统一至正文、标题和Figures1/4/5图内；Figure2用adjusted-association；4.2按要求使用Student and Contextual Characteristics。','- 3.3含预测、SHAP、DML三个分节及原生公式；Figure1仍在第二章。Figure1/2/4与Table2标题同步修订。','- 2章解释概念，3章记录操作，4章报告发现，5章呈现应用和教师结果，6章教育解释，7章集中限制。5.4保留简短评价目的并回指3.5/3.6。','','## 2 原始数据、方法和数字','','- 原始SAV低内存分段核查成功：日本5760人，182所学校，十个数学合理值均观测；种子42划分4608/1152。42个预测特征在训练集中均未超过50%缺失。','- 五模型×十合理值50行指标、42特征×十合理值420行SHAP、12因素×十合理值120行DML均从明细重算，与汇总一致。','- 原稿方法描述实际错误已改：树模型训练交叉验证是3折；ElasticNet内部5折。MATH_MEAN参与训练调参，最后模型按十合理值分别拟合。','- DML初始14因素，FEELSAFE和REPEAT在日本样本全缺失，筛选后12。原DML变量池21列，每个因素由其余20列调整，不是扩展预测模型42列；CNT为常数，不作为国家比较。','- 数值中位数/分类众数插补规则及折内处理按原代码描述。SHAP均匀抽取至多3000人，随后按学生权重汇总，不是按权重概率抽样。','- 学校聚类标准误及合理值合并方差Ubar+(1+1/10)B核查通过；没有声称使用未执行的80复制权重完整推断。','- 12案例×5次×2系统=120条回复；两教师评分=240行、960个维度分数。均值、配对差值、区间及二次加权kappa从原评分重新核对。Bootstrap为50000次、按12配对案例重采样、seed42；原生成/匿名顺序种子不改。','','## 3 教育主线与推断范围','','实证分析优先，智能体用于呈现、讨论和建议。本文可支持教师对这些案例最终回复的质量评价，不证明真实学生诊断准确、教师决策改善或学生成绩提高，也不能将整套系统差异归因于单个组件。负向教师支持关联按OECD原文解释为可能存在补偿性支持；保留其他解释及时间顺序限制。','','## 4 逐条引文与书目','','26条正文与文后双向对应，未检出缺项或无正文引用条目。23个DOI对应的标题、作者、年份、刊名、卷期、页码/文章号逐项检查；两条NeurIPS原会议记录，Hu依作者提供。以下明确区分摘要和全文核查，不声称26篇全部全文精读。','','|条目|核查层级|具体结论|一手依据|','|---|---|---|---|']
for r in rows:report.append('|'+r['reference']+'|'+r['scope']+'|'+r['assessment']+'|'+r['source']+'|')
report+=['','书目修正：补Black、Bond、Chernozhukov、Kuhail、Shmueli、Yan和Zawacki-Richter的期号；Khosravi第二作者改为Buckingham Shum, S.，Tsai和Lee名缩写加连字符；Alnasyan第三作者按出版者姓名写M.；AI期刊全名用and；Albreiki题名按出版者元数据保留其不寻常的student’拼写，而非自行修正文献标题。Lundberg补原会议链接。','','## 5 字数、图表与投稿条件','',f"摘要{counts['abstract_words_space_delimited']}词（空格分词）；正文含标题和图注、不含摘要/表格/参考文献为{counts['main_words_including_headings_and_captions_excluding_abstract_references_tables']}词；全文含摘要/表格/参考文献为{counts['all_words_including_abstract_references_tables']}词。计数口径与Word可能略异。",'','官方SEE指南及Elsevier备用指南入口均403，Editorial Manager公共入口也未能读取；不以第三方9000词或150词数字冒充最新官方要求。简明摘要属于编辑处理，不能宣称已确认期刊全部格式。','Supplementary Methods的S10映射已在新ZIP显式补齐；统计代码、结果与图使用seed42版本，原始回复和教师记录不变。','', '## 6 仅剩的明确待确认项','','1. Hu文献接收/在印状态需要作者实际出版或接收材料；不补造DOI。','2. 最新官方作者指南仍需可读取页面/导出文件，才能最终判定字数、关键词及投稿附件要求。','3. 图3按作者要求保留旧界面，需作者确认其能代表本轮评价所用交互。','4. 匿名评分不能等同于伦理审批、豁免或知情同意；这些事实尚未获确认，未虚构伦理声明。','5. 作者、通讯、贡献、项目及利益冲突按作者要求暂缓。','', '以上不阻止已完成的正文、结果与引文核查，但不据此标记整篇已满足所有投稿条件。']
report.insert(report.index('## 3 教育主线与推断范围'),'公式记号核查另修正一处：式(1)结构函数g(X)与残差化所需的条件均值E(Y|X)不应混同。式(3)及说明改用ℓ(X)=E(Y|X)，保留m(X)=E(D|X)。仅修正原生Word公式记号，原计算代码和结果未变。')
(F/'GRADE-Agent_核查与修订记录_v35.md').write_text('\n'.join(report),encoding='utf-8')
with zipfile.ZipFile(F/'Supplementary_Evaluation_Materials_v34.zip')as z:items={n:z.read(n)for n in z.namelist()}
items['audit/raw_sample_audit.json']=(BASE/'outputs/submission_audit_v33_20261028/raw_sample_audit.json').read_bytes()
items['audit/numeric_audit.json']=(BASE/'outputs/submission_audit_v33_20261028/numeric_audit.json').read_bytes()
items['SHA256SUMS.json']=json.dumps({n:hashlib.sha256(v).hexdigest()for n,v in items.items()if n!='SHA256SUMS.json'},indent=2).encode()
with zipfile.ZipFile(F/'Supplementary_Evaluation_Materials_v35.zip','w',zipfile.ZIP_DEFLATED)as z:
 for n,v in items.items():z.writestr(n,v)
print(json.dumps(counts,indent=2));print('Reference/caption/formula/table checks passed. Rendering.',flush=True)
renderer=Path('C:/Users/84130/.codex/plugins/cache/openai-primary-runtime/documents/26.1025.11240/skills/documents/render_docx.py')
subprocess.run([sys.executable,str(renderer),str(target),'--output_dir',str(OUT/'render'),'--emit_pdf'],check=True)

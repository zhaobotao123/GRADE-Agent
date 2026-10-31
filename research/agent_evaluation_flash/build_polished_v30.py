from pathlib import Path
from copy import deepcopy
from collections import Counter
import re,json,hashlib
from zipfile import ZipFile
from docx import Document
from docx.shared import Pt,Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from polish_v30_text import REVISIONS

FOLDER=Path('D:/Users/84130/Desktop/投稿/转投修改_20261022')
SOURCE=FOLDER/'GRADE-Agent_SEE_Teacher_Evaluation_v29.docx'
TARGET=FOLDER/'GRADE-Agent_SEE_NatureSkills_Polished_v30.docx'
QA=Path('outputs/manuscript_polished_v30_20261027')
QA.mkdir(parents=True,exist_ok=True)
doc=Document(SOURCE)
old=list(doc.paragraphs)
original=[p._p.xml for p in old]
oldmath=[__import__('lxml').etree.tostring(el) for el in doc._element.xpath('.//m:oMath')]
tables=[t._tbl.xml for t in doc.tables]
oldref=next(i for i,p in enumerate(old) if p.text=='References')
audit=[]
numeric=lambda s:Counter(re.findall(r'(?<![A-Za-z])[-−]?\d+(?:[.,]\d+)*(?:%|²)?',s))
for i,new in sorted(REVISIONS.items()):
    p=old[i]
    assert not p._p.xpath('.//m:oMath|.//w:drawing'), (i,'protected object')
    before=p.text
    rpr=deepcopy(p.runs[0]._r.rPr) if p.runs else None
    p.clear()
    r=p.add_run(new)
    if rpr is not None:r._r.insert(0,rpr)
    if 8<=i<205:
        p.paragraph_format.first_line_indent=Inches(.25)
        p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    audit.append({'paragraph':i,'before':before,'after':new,'numeric_tokens_unchanged':numeric(before)==numeric(new)})
    if i in (105,107,108,110,111,113,156,158,160,161,203):
        assert numeric(before)==numeric(new), (i,numeric(before)-numeric(new),numeric(new)-numeric(before))

# Keep the abstract together without altering the already selected title or keywords.
for i in range(2,6):
    old[i].paragraph_format.line_spacing=1.35
    old[i].paragraph_format.keep_with_next=True
old[6].paragraph_format.keep_with_next=False
old[7].paragraph_format.page_break_before=True

assert tables==[t._tbl.xml for t in doc.tables]
assert oldmath==[__import__('lxml').etree.tostring(el) for el in doc._element.xpath('.//m:oMath')]
assert [p._p.xml for p in old[oldref:]]==original[oldref:]
for i,p in enumerate(old):
    if i not in REVISIONS and i not in range(2,8):assert p._p.xml==original[i], i
doc.save(TARGET)
with ZipFile(SOURCE) as a, ZipFile(TARGET) as b:
    for name in a.namelist():
        if name.startswith('word/media/'):assert a.read(name)==b.read(name),name
count=lambda d:sum(len(p.text.split()) for p in d.paragraphs)+sum(len(c.text.split()) for t in d.tables for row in t.rows for c in row.cells)
src=Document(SOURCE)
new=Document(TARGET)
summary={'source':str(SOURCE),'target':str(TARGET),'changed_paragraphs':len(audit),'words_including_tables_and_references_before':count(src),'words_including_tables_and_references_after':count(new),'abstract_words':sum(len(p.text.split()) for p in new.paragraphs[2:6]),'tables_unchanged':True,'equations_unchanged':True,'embedded_images_unchanged':True,'reference_list_unchanged':True,'result_numeric_tokens_verified':True,'changes':audit}
(QA/'revision_audit.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')

notes=f'''# 全文润色说明

本次以 v29 为基础，使用 nature-polishing 与 paper-writing 的章节分工、去重和证据边界检查；表达面向 Studies in Educational Evaluation 的教育研究读者，而非套用 Nature 的技术叙事。

## 已完成的修改

- 逐段检查摘要、引言、概念框架、方法、实证结果、智能体设计与评价、讨论、局限、结论及可用性声明；其中 {len(audit)} 段进行了语言修改。原题目、关键词和编号层次保留。
- 减少反复出现的研究目的、人口层面与个体层面区别，以及平台价值的重复说明；保留原有主张和必要限定。
- 引言继续以教育问题切入，保持“成绩相关因素分析在先，智能体应用与评价在后”的贡献顺序。
- 方法部分保留候选变量、缺失值阈值、训练测试划分、合理值、权重、交叉拟合、评分锚点和自助法细节。公式及含公式段落保持原样。
- 统一英式拼写，以及 adjusted associations、predictive contributions、ordinary multi-turn LLM 等用语。评分维度在首次定义与图表中保留全称。
- 表格数值、方程、嵌入图片和参考文献表未改；结果段落的数字标记逐项核对。未重新运行实验、改评分或生成新研究结果。

## 篇幅统计

- 润色前：约 {count(src):,} 词。
- 润色后：约 {count(new):,} 词。
- 摘要：约 {summary['abstract_words']} 词。

以上为段落及表格按空白分词统计，包含参考文献和图表文字，不包含图片内文字；Word 与投稿系统的统计可能略有差异。不能据此宣称已经满足期刊最新篇幅要求。

## 需要作者确认的实质问题

1. **完整系统比较的解释**：6.3 明确补充“不能将差异归因于单个组件”，与现有方法、局限中的边界一致。未削弱或改变系统评分更高这一结果。
2. **DML 因素选择的表述**：将引言中可能暗示“根据预测或 SHAP 结果选择 DML 因素”的措辞，统一为“预先指定因素的独立分析”，与原方法一致。这是消除内部歧义，不是改变实验。
3. **教师评价伦理与同意**：当前材料没有确认的审批／豁免／知情同意信息，本次没有编写声明。请依据真实安排补充。
4. **图3的对应关系**：按你的选择保留原图和图注。旧界面截图含报告式展示，而本次评价对象是对话后的简短回复；请在正式投稿前确认是否需要图注说明，或更换成当前界面。
5. **方法可复现性**：文中未具体说明类别缺失值的填补规则、完整超参数与软件版本；补充材料需确认可定位这些信息。正文未补造参数，也未修改冻结实验代码。
6. **投稿附件**：作者信息页、Highlights、利益冲突、资助及作者贡献等仍需按实际情况和当前投稿系统要求单独准备。
7. **引文核查范围**：本次保留原有引文与文献条目，没有新增文献，也不等同于逐篇重新核验文献内容、DOI或期刊最新作者指南。

## 术语对应

|英文用语|中文含义|本次处理|
|---|---|---|
|mathematics achievement|数学成绩|用于明确本研究的分析结果变量|
|predictive contribution|预测贡献|不等同于因果影响|
|adjusted association|调整后关联|不改写为已证实的因果效应|
|learning-support suggestions|学习支持建议|不改写为已验证的干预效果|
|teacher-rated reply quality|教师评价的回复质量|与真实学生的诊断准确性、学习效果区分|
|inter-rater agreement|评分者一致性|与两系统平均分差异分别报告|

本次交付的是语言与结构润色稿，不是期刊接收保证或完整方法学验证。原版本与原始数据均保留。
'''
(FOLDER/'GRADE-Agent_全文润色说明_v30.md').write_text(notes,encoding='utf-8')
print(json.dumps({k:v for k,v in summary.items() if k!='changes'},ensure_ascii=True,indent=2))

"""Read-only numeric, bibliographic and protocol checks for the submission revision."""
import os
os.environ['OPENBLAS_NUM_THREADS'] = '1'
os.environ['OMP_NUM_THREADS'] = '1'
import sys, csv, json, re, hashlib, urllib.request, urllib.parse, time
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor
from docx import Document
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'work/figure_python_deps'))
import numpy as np

BASE = Path(__file__).resolve().parents[1]
OUT = BASE / 'outputs/submission_audit_v33_20261028'
OUT.mkdir(parents=True, exist_ok=True)
FOLDER = Path('D:/Users/84130/Desktop/投稿/转投修改_20261022')
SOURCE = FOLDER / 'GRADE-Agent_SEE_Restructured_v32.docx'

def readcsv(path):
    with Path(path).open(encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f))

def write(name, obj):
    (OUT/name).write_text(json.dumps(obj, ensure_ascii=False, indent=2), encoding='utf-8')

def hashfile(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()

def numeric_checks():
    formal = Path('D:/pisa2022/outputs_grade_agent_formal')
    tuned = Path('D:/pisa2022/outputs_five_models_tuned_42_features')
    shap = Path('D:/pisa2022/outputs_final_shap_42_features')
    checks = {}
    modelpv = readcsv(tuned/'tuned_five_model_metrics_by_pv.csv')
    models = readcsv(tuned/'tuned_five_model_metrics_summary.csv')
    assert len(modelpv) == 50 and len(models) == 5
    for row in models:
        rows = [r for r in modelpv if r['model'] == row['model']]
        assert len(rows) == 10
        for metric in ('weighted_r2', 'weighted_rmse', 'weighted_mae'):
            np.testing.assert_allclose(np.mean([float(r[metric]) for r in rows]), float(row[metric+'_mean']), atol=1e-12)
            np.testing.assert_allclose(np.std([float(r[metric]) for r in rows], ddof=1), float(row[metric+'_sd']), atol=1e-12)
    checks['five_models_ten_pvs'] = {'pass': True, 'metrics': models}
    cv = readcsv(tuned/'training_cv_all_candidates.csv')
    checks['outer_training_cv_folds'] = sorted({r['fold'] for r in cv})
    assert checks['outer_training_cv_folds'] == ['1', '2', '3']
    dmlpv = readcsv(formal/'dml_results_by_pv.csv')
    dml = readcsv(formal/'dml_results_pooled.csv')
    assert len(dmlpv) == 120 and len(dml) == 12
    for pooled in dml:
        rows = [r for r in dmlpv if r['treatment'] == pooled['treatment']]
        assert len(rows) == 10
        estimates = np.array([float(r['coef']) for r in rows])
        se = np.array([float(r['std_error']) for r in rows])
        total_var = np.mean(se**2) + 1.1*np.var(estimates, ddof=1)
        np.testing.assert_allclose(estimates.mean(), float(pooled['coef_pooled']), atol=1e-12)
        np.testing.assert_allclose(total_var**.5, float(pooled['std_error_pooled']), atol=1e-12)
        assert all(r['inference'] == 'school_clustered_CNTSCHID' for r in rows)
        assert all(r['dml_method'] == 'weighted_cross_fitted_plr_3fold' for r in rows)
    checks['dml_pooling_recomputed_from_120_rows'] = {'pass': True, 'pooled': dml}
    shappv = readcsv(shap/'shap_importance_by_pv.csv')
    shaps = readcsv(shap/'shap_importance_pooled.csv')
    assert len(shaps) == 42 and len(shappv) == 420
    for row in shaps:
        rows = [r for r in shappv if r['feature'] == row['feature']]
        assert len(rows) == 10
        np.testing.assert_allclose(np.mean([float(r['mean_abs_shap']) for r in rows]),float(row['mean_abs_shap']),atol=1e-12)
    checks['42_original_predictors_and_shap_ten_pvs'] = {'pass': True, 'features': [r['feature'] for r in shaps]}
    teachers = BASE/'outputs/manuscript_teacher_evaluation_v21_20261027'
    rows = readcsv(teachers/'teacher_scores_long.csv')
    dims = ['Contextual relevance','Consistency with case facts','Appropriateness and actionability','Clarity']
    cases = sorted({r['case_id'] for r in rows})
    jobs = {r['source_job'] for r in rows}
    assert len(rows)==240 and len(jobs)==120 and len(cases)==12
    assert {r['teacher'] for r in rows} == {'1','2'}
    for teacher in ('1','2'):
        for case in cases:
            for system in ('ordinary_llm','grade_agent'):
                rr=[r for r in rows if r['teacher']==teacher and r['case_id']==case and r['system']==system]
                assert len(rr)==5
                assert all(float(r[d]) in (1,2,3,4,5) for r in rr for d in dims)
    cube=np.array([[[np.mean([float(r[d]) for r in rows if r['case_id']==c and r['system']==s]) for d in dims]
                    for s in ('ordinary_llm','grade_agent')] for c in cases])
    delta=cube[:,1]-cube[:,0]
    draws=np.random.default_rng(42).integers(0,12,size=(50000,12))
    ci=np.quantile(delta[draws].mean(axis=1),[.025,.975],axis=0)
    matched={(r['teacher'],r['source_job']):r for r in rows}
    frozen=json.loads((BASE/'outputs/teacher_bootstrap_seed42_20261028/results_seed42.json').read_text(encoding='utf-8'))
    results=[]
    for j,d in enumerate(dims):
        obs=np.zeros((5,5))
        for k in sorted(jobs): obs[int(matched['1',k][d])-1,int(matched['2',k][d])-1]+=1
        expected=np.outer(obs.sum(axis=1),obs.sum(axis=0))/120
        weights=(np.arange(5)[:,None]-np.arange(5)[None,:])**2
        result=dict(dimension=d, ordinary_mean=float(cube[:,0,j].mean()), agent_mean=float(cube[:,1,j].mean()),
                    paired_difference=float(delta[:,j].mean()), ci_low=float(ci[0,j]),ci_high=float(ci[1,j]),
                    quadratic_weighted_kappa=float(1-(weights*obs).sum()/(weights*expected).sum()))
        for key in result:
            if key!='dimension': np.testing.assert_allclose(result[key],frozen['summary'][j][key],atol=1e-12)
        results.append(result)
    checks['teacher_scores_independently_recomputed']={'pass':True,'score_count':960,'reply_count':120,'case_count':12,'runs_per_system_case':5,'teacher_count':2,'bootstrap_seed':42,'bootstrap_draws':50000,'results':results}
    checks['source_hashes']={str(p):hashfile(p) for p in [formal/'dml_results_by_pv.csv',formal/'dml_results_pooled.csv',tuned/'tuned_five_model_metrics_by_pv.csv',shap/'shap_importance_by_pv.csv',teachers/'teacher_scores_long.csv']}
    write('numeric_audit.json',checks)
    print('Numerical checks passed; 3 outer training CV folds confirmed.',flush=True)

def fetch_crossref(ref):
    if not ref.get('doi'):
        return {'status':'manual_needed','reason':'No DOI in supplied reference; use primary conference/publisher record.'}
    url='https://api.crossref.org/works/'+urllib.parse.quote(ref['doi'],safe='')
    for attempt in range(3):
        try:
            req=urllib.request.Request(url,headers={'User-Agent':'GRADEAgentReferenceCheck/1.1 (read-only scholarly metadata audit)'})
            with urllib.request.urlopen(req,timeout=22) as response: m=json.load(response)['message']
            metadata={k:m.get(k) for k in ['DOI','title','container-title','author','editor','volume','issue','page','article-number','published','published-print','published-online','issued','abstract','link','URL','type']}
            return {'status':'metadata_retrieved','url':url,'metadata':metadata}
        except Exception as exc:
            reason=type(exc).__name__+': '+str(exc)
            if attempt<2: time.sleep(1.5*(attempt+1))
    return {'status':'manual_needed','url':url,'reason':reason}

def reference_checks():
    doc=Document(SOURCE)
    ri=next(i for i,p in enumerate(doc.paragraphs) if p.text=='References')
    body='\n'.join(p.text for p in doc.paragraphs[:ri])+'\n'+'\n'.join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    refs=[]
    for p in doc.paragraphs[ri+1:]:
        if not p.text.strip():continue
        text=p.text
        doi=re.search(r'https://doi.org/([^\s]+)',text)
        refs.append({'entry':text,'first_author':text.split(',')[0] if not text.startswith('OECD.') else 'OECD',
                     'year':re.search(r'\((\d{4}[ab]?|in press)\)',text).group(1),
                     'doi':doi.group(1).rstrip('.') if doi else None})
    with ThreadPoolExecutor(max_workers=2) as executor:
        for ref,remote in zip(refs,executor.map(fetch_crossref,refs)):
            ref['verification']=remote
            print(ref['first_author'],ref['year'],remote['status'],flush=True)
    write('reference_metadata.json',{'references':refs,'main_citation_text':body,'scope':'DOI metadata including complete author list; claim support separately audited against available primary text.'})

if __name__=='__main__':
    numeric_checks()
    reference_checks()

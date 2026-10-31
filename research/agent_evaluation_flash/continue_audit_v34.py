import os,sys,json,re,textwrap,io,csv,zipfile,hashlib,subprocess
from pathlib import Path
from copy import deepcopy
os.environ['OPENBLAS_NUM_THREADS']='1';os.environ['OMP_NUM_THREADS']='1'
BASE=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(BASE/'work/figure_python_deps'))
from docx import Document
from docx.oxml.ns import qn
from lxml import etree
from PIL import Image
import finalize_teacher_results_v21 as plotting
F=Path('D:/Users/84130/Desktop/投稿/转投修改_20261022')
OUT=BASE/'outputs/submission_audit_v34_20261028';OUT.mkdir(parents=True,exist_ok=True)
figures=OUT/'figures';figures.mkdir(exist_ok=True)
results=json.loads((BASE/'outputs/teacher_bootstrap_seed42_20261028/results_seed42.json').read_text(encoding='utf-8'))
source=(BASE/'agent_evaluation_flash/finalize_teacher_results_v21.py').read_text(encoding='utf-8')
start=source.index(" font_manager.findfont('Times New Roman'")
end=source.index(" save(OUT/'qa/figure_contract.json'")
plotcode=textwrap.dedent(source[start:end]).replace('Ordinary multi-turn LLM','Ordinary LLM')
ns=vars(plotting).copy();ns.update(OUT=OUT,summary=results['summary'])
exec(compile(plotcode,'audited_figure_reexport','exec'),ns)
d=Document(F/'GRADE-Agent_SEE_Audited_Draft_v33.docx')
math=[etree.tostring(x)for x in d._element.xpath('.//m:oMath')]
tables=[t._tbl.xml for t in d.tables]
def replace(p,t):
    assert not p._p.xpath('.//m:oMath|.//w:drawing')
    prop=deepcopy(p.runs[0]._r.rPr)if p.runs else None
    p.clear();r=p.add_run(t)
    if prop is not None:r._r.insert(0,prop)
abstract="Understanding achievement-related factors requires both empirical analysis and accessible educational interpretation. Using Japan’s PISA 2022 sample, we combined weighted predictive modelling, SHAP feature attribution, and Double Machine Learning to examine mathematics achievement. Mathematics self-efficacy and disciplinary climate showed stable positive adjusted associations, whereas support-related findings required contextual interpretation. We developed GRADE-Agent to discuss these findings through conversation and brief learning-support summaries. Two teachers independently rated 120 English replies from GRADE-Agent and an ordinary large language model across twelve synthetic cases and five runs per system. GRADE-Agent received higher mean ratings on all four criteria. The largest difference concerned suggestion appropriateness and actionability (4.75 versus 3.61; paired difference 1.14, 95% case-cluster bootstrap confidence interval [0.86, 1.45]). The findings support further evaluation of research-informed conversational support, without establishing individual diagnostic accuracy or improved educational outcomes."
absidx=next(i for i,p in enumerate(d.paragraphs)if p.text=='Abstract')
paras=list(d.paragraphs);replace(paras[absidx+1],abstract)
for p in paras[absidx+2:]:
    if p.text.startswith('Keywords:'):break
    p._p.getparent().remove(p._p)
for p in d.paragraphs:
    t=p.text.replace('12 prespecified factors','12 retained candidate factors').replace('12 pre-specified candidate factors','12 retained candidate factors')
    t=t.replace('Supplementary Section S10','Supplementary Methods, Section S10')
    if t!=p.text and not p._p.xpath('.//m:oMath|.//w:drawing'):replace(p,t)
for index,name in [(0,'Figure1_updated_framework'),(3,'Figure4_evaluation_protocol'),(4,'Figure5_teacher_comparison')]:
    shape=d.inline_shapes[index];blip=shape._inline.graphic.graphicData.pic.blipFill.blip
    png=(figures/(name+'.png')).read_bytes();d.part.related_parts[blip.embed]._blob=png
    im=Image.open(io.BytesIO(png));shape.height=int(shape.width*im.height/im.width)
    for elem in shape._inline.iter():
        if etree.QName(elem).localname=='svgBlip':d.part.related_parts[elem.get(qn('r:embed'))]._blob=(figures/(name+'.svg')).read_bytes()
assert math==[etree.tostring(x)for x in d._element.xpath('.//m:oMath')]
assert tables==[t._tbl.xml for t in d.tables]
target=F/'GRADE-Agent_SEE_Checked_v34.docx';d.save(target)
metadata=json.loads((BASE/'outputs/submission_audit_v33_20261028/reference_metadata.json').read_text(encoding='utf-8'))
ri=next(i for i,p in enumerate(d.paragraphs)if p.text=='References')
body='\n'.join(p.text for p in d.paragraphs[:ri])
hits=[]
for ref in metadata['references']:
    author=ref['first_author'];year=ref['year']
    pattern=re.escape(author)+r'[^;\n]{0,70}?'+re.escape(year)
    matches=re.findall(pattern,body)
    if author=='OECD' and year=='2023b' and 'OECD, 2023a, 2023b' in body:matches.append('OECD, 2023a, 2023b')
    hits.append((author,year,bool(matches)))
report=['# 前四项核查与投稿准备记录 v34','',
'## 已完成的文字与图表处理','',
'正文统一ordinary LLM；Figure 4内部同名标签已从Python绘图源修正；Figure 1、4、5重新导出，Figure 5使用种子42复算数据。Figure 2内部adjusted-association已与正文一致，本轮未改变原始数据。原生公式、教师分数和表格保持不变。摘要压缩为'+str(len(abstract.split()))+'个空格分词；这属于编辑选择，不将尚未核实的150词上限说成当前期刊规定。', '',
'第三章说明实际研究流程，第七章说明这些流程限制了哪些推断；相同案例数用于定义设计及其适用范围，不再将二者误作两个独立实验。3.3保留三种分析方法的分节和已有DML原生公式。', '',
'## 方法和数字核查','',
'详见上一轮原始分析审计：50行预测结果、420行SHAP、120行DML和240行教师评分记录重算一致；共120条回复、960个评分。树模型调参实际三折，ElasticNet内部五折；合理值均值用于调参；DML初始14因素筛选后12因素，调整池21变量而非扩展预测42变量。原始SAV读取内存失败，5760仍依据原日志与描述表，未冒称新一次原始数据验证。','',
'## 参考文献逐条对应记录','',
'以下为逐条审计状态，不把元数据通过当作全文支持通过。23条DOI已通过CrossRef取得对应记录；Ke和Lundberg通过NeurIPS原始会议记录确认。Hu的in-press状态依作者提供，尚未独立确认。','',
'|条目|文内检出|元数据依据|论述核查范围|','|---|---|---|---|']
available={'Albreiki':'CrossRef所载作者摘要支持学生表现预测的一般背景；未验证特定中学生适用性。','Bond':'出版者摘要支持伦理、情境及跨学科需求，不能据此声称GRADE-Agent已有效。','Fabbri':'ACL原始页面支持文本生成评价背景；本研究采用改编量表，不称其为经验证教育量表。','Liu':'ACL原始页面确认G-Eval是LLM评价方法；本研究为人工评分，正文明确未实施G-Eval。','Kuhail':'出版者摘要支持教育聊天机器人及互动设计背景，不证明本系统追问有效。','Shmueli':'作者摘要支持预测与解释的区分。','Zawacki-Richter':'出版者摘要支持教育视角与伦理研究缺口。','Wang':'2026作者机构摘要确认57项研究；2023文章全文仍受限，相关具体解释需进一步核查。','Jin':'作者机构摘要支持在线同伴评价行为与预测研究的描述。','Lee':'出版者摘要和方法节选支持AI辅助评价及学生反馈的描述。','OECD':'对应出版物页面确认年份与书目；2023a Volume I、2023b Volume II一致。2010教师补偿性支持的具体页码仍需全文定位。','Ke':'NeurIPS原始摘要确认LightGBM及全部8位作者。','Lundberg':'NeurIPS原始摘要确认SHAP贡献解释及两位作者。','Black':'出版者摘要支持形成性评价理论方向；具体定义的全文页码尚需定位。','Hu':'作者提供的手稿/书目信息，未独立证实接收及出版状态。'}
for ref,hit in zip(metadata['references'],hits):
    a=ref['first_author'];year=ref['year'];remote=ref['verification']
    url='https://doi.org/'+ref['doi']if ref.get('doi')else ('https://proceedings.neurips.cc/'if a in('Ke','Lundberg')else '作者提供')
    report.append('|'+a+' '+year+'|'+('有'if hit[2]else '未检出，需人工复核')+'|'+url+'|'+available.get(a,'DOI元数据对应；本轮访问全文失败/受限，尚不能确认具体论述充分支持。')+'|')
report+=['','## 期刊要求','',
'官方作者指南仍403；已尝试官方链接及搜索索引，不能以第三方9000词/150词限制代替最新官方核实。当前正文和全文计数约6300/7800词，精确计数随摘要调整变化；投稿前应取得可访问的官方指南或投稿系统说明。','',
'## 尚需外部材料的具体事项','',
'作者要求暂缓作者/贡献和经费，不填假信息；匿名评分不等于已获伦理豁免或知情同意。待确认项不阻止其余审计，但也不标记整篇已可直接投稿。图3保留作者选择的旧界面图，仍需说明其与本次实际评价系统的对应。部分全文访问限制及Hu出版状态需作者提供原文/接收材料后完成。']
(F/'GRADE-Agent_前四项核查记录_v34.md').write_text('\n'.join(report),encoding='utf-8')
# Build the referenced factor mapping explicitly within the current package.
with zipfile.ZipFile(F/'Supplementary_Evaluation_Materials_v33_seed42.zip')as z:items={n:z.read(n)for n in z.namelist()}
with Path('D:/pisa2022/outputs_grade_agent_formal/dml_results_pooled.csv').open(encoding='utf-8-sig')as f:rows=list(csv.DictReader(f))
supp=['# Supplementary Methods','','## S10. Factor-level interpretation mapping','','|Factor|Interpretation class|Scope|','|---|---|---|']
for r in rows:
    factor=r['treatment'];cl=('Stable positive'if factor in('MATHEFF','DISCLIM')else 'Boundary review'if factor in('BELONG','ICTSCH','TEACHSUP')else 'Directionally unstable'if factor in('ANXMAT','MATHEF21')else 'Weak or context-sensitive')
    supp.append(f'|{factor}|{cl}|Population-level adjusted association; not an established individual cause.|')
items['Supplementary_Methods.md']='\n'.join(supp).encode()
for file in figures.glob('*'):
    if file.suffix in('.png','.svg','.pdf'):items['figures/'+file.name]=file.read_bytes()
items['README.md']+=b'\nSupplementary_Methods.md contains Section S10, the factor-level interpretation mapping cited by the manuscript. Updated figures use Ordinary LLM terminology.\n'
items['SHA256SUMS.json']=json.dumps({n:hashlib.sha256(v).hexdigest()for n,v in items.items()if n!='SHA256SUMS.json'},indent=2).encode()
with zipfile.ZipFile(F/'Supplementary_Evaluation_Materials_v34.zip','w',zipfile.ZIP_DEFLATED)as z:
    for n,v in items.items():z.writestr(n,v)
print(target);print('Abstract words',len(abstract.split()));print('Unmatched reference entries',[x for x in hits if not x[2]])
renderer=Path('C:/Users/84130/.codex/plugins/cache/openai-primary-runtime/documents/26.1025.11240/skills/documents/render_docx.py')
subprocess.run([sys.executable,str(renderer),str(target),'--output_dir',str(OUT/'render'),'--emit_pdf'],check=True)

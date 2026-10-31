import json,re
from pathlib import Path
from docx import Document
F=Path('D:/Users/84130/Desktop/投稿/转投修改_20261022')
d=Document(F/'GRADE-Agent_SEE_Checked_v34.docx')
ri=next(i for i,p in enumerate(d.paragraphs)if p.text=='References')
refs=json.loads(Path('outputs/submission_audit_v33_20261028/reference_metadata.json').read_text(encoding='utf-8'))['references']
for r in refs:
 m=r['verification'].get('metadata',{})
 print(r['entry'])
 print('REMOTE',m.get('volume'),m.get('issue'),m.get('page'),m.get('article-number'),'; '.join(a.get('family','')+', '+a.get('given','')for a in m.get('author',[])))
for p in d.paragraphs[:ri]:
 if re.search(r'\d{4}[ab]?\)|in press\)',p.text):print('CITE',p.text)

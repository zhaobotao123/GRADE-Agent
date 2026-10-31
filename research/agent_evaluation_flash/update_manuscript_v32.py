"""Apply author-approved heading architecture while preserving results and figures."""
from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile
import hashlib
import json
import re

from docx import Document
from lxml import etree

FOLDER = Path('D:/Users/84130/Desktop/投稿/转投修改_20261022')
SOURCE = FOLDER / 'GRADE-Agent_SEE_Refined_v31.docx'
TARGET = FOLDER / 'GRADE-Agent_SEE_Restructured_v32.docx'
QA = Path('outputs/manuscript_structure_v32_20261028')
QA.mkdir(parents=True, exist_ok=True)
d = Document(SOURCE)
old = list(d.paragraphs)
tables = [t._tbl.xml for t in d.tables]
math = [etree.tostring(x) for x in d._element.xpath('.//m:oMath')]

HEADINGS = {
7: '1. Introduction',
23: '2. Conceptual Framework',
24: '2.1 Predictive Contributions and Adjusted Associations',
28: '2.2 Formative Assessment and Educational Support',
33: '2.3 Connecting Research Findings with Students’ Learning Situations',
40: '3. Methods',
42: '3.1 Educational Context and Data Source',
48: '3.2 Achievement-Related Factors and Data Preparation',
66: '3.3 Empirical Analysis of Achievement-Related Factors',
83: '3.4 Translating Empirical Findings into Educational Information',
88: '3.5 Case Development and Teacher Evaluation Criteria',
96: '3.6 System Comparison and Statistical Analysis',
102: '4. Results',
103: '4.1 Predictive Model Performance',
105: '4.2 Predictive Contributions of Candidate Factors',
108: '4.3 Adjusted Associations between Candidate Factors and Achievement',
114: '4.4 Summary of Findings Informing GRADE-Agent',
119: '5. GRADE-Agent: Translating Achievement-Related Findings into Educational Support',
120: '5.1 Connecting Achievement-Related Findings with Educational Questions',
125: '5.2 Exploring Students’ Learning Situations through Conversation',
130: '5.3 The GRADE-Agent Interface and Learning-Support Summaries',
136: '5.4 Teacher Evaluation Procedure',
150: '5.5 Teacher Ratings and Implications for Educational Support',
166: '6. Discussion',
167: '6.1 Interpreting Achievement-Related Findings for Educational Support',
174: '6.2 Interpreting Support-Related Associations in Context',
181: '6.3 Comparing GRADE-Agent with an Ordinary LLM',
186: '6.4 Implications for Educational Use and Future Evaluation',
192: '7. Limitations and Future Research',
198: '8. Conclusion',
}
EDITS = {
35: "Conversation allows users to clarify circumstances and revise initial interpretations (Chiu et al., 2023; Kuhail et al., 2023; Yan et al., 2024). This process brings research findings into discussion alongside observations of the learner's tasks, experiences, and existing support. Explanations remain provisional while relevant information is missing. Educators judge whether the interpretation fits the learner and which next steps are appropriate.",
36: "An explanation's research basis and its educational usefulness require separate evaluation. Relevant statistical findings do not ensure faithful interpretation of a case or understandable, practical advice. Evaluation therefore needs to consider how well a reply relates findings to the learning situation and supports professional judgement.",
84: "We organised predictive, SHAP, and DML findings in an empirical evidence knowledge base for GRADE-Agent. Each factor record included its predictive contribution, DML estimate where available, uncertainty, and interpretive limits. Records distinguished predictive information from adjusted associations and retained their population-level scope.",
86: "The application used these records to identify areas for inquiry, request missing contextual information, and discuss possible learning-support options. Context supplied by the user determined which findings were relevant to the conversation. Section 5 describes their presentation through educational questions, dialogue, and brief summaries.",
95: "Clarity assessed whether teachers or parents could follow the reply. A rating of 1 indicated difficult expression or organisation; 3 indicated understandable writing with unnecessary terminology, repetition, or organisational problems. A rating of 5 indicated clear, coherent, concise wording. Contextual relevance, case-fact consistency, and clarity were secondary outcomes. Sections 5.4 and 5.5 summarise the procedure and present the paired results.",
99: old[99].text.replace('seed 20261027', 'seed 42'),
111: "We assigned BELONG, ICTSCH, and TEACHSUP to boundary review because their negative adjusted associations required contextual examination. Section 6.2 considers their educational interpretation, including the timing of assistance and the distinction between resource availability and use.",
121: "GRADE-Agent links each relevant finding to questions that educators can explore within the described learning situation. For mathematics self-efficacy, discussion focuses on confidence and approaches to mathematical tasks. Users can add observations about independent work, persistence, and responses to assistance before considering possible support.",
123: "Teacher support requires a different line of inquiry because its adjusted association was negative. The conversation asks when additional help began and which difficulties preceded it. This question helps distinguish an existing learning need from an unsupported conclusion about the effects of assistance. Research informs the question, while the learner's circumstances guide the interpretation.",
138: "Two teachers independently evaluated anonymous final replies from GRADE-Agent and an ordinary multi-turn LLM against the full case descriptions. The comparison included 120 English final replies across 12 synthetic cases, with five runs per system and identical staged information.",
139: "The four dimensions addressed contextual relevance, consistency with case facts, suggestion appropriateness and actionability, and clarity. Together, they examined whether replies attended to the learning situation, used information faithfully, offered suitable next steps, and communicated understandably.",
145: "Suggestion appropriateness and actionability was the prespecified primary outcome. Sections 3.5 and 3.6 provide the case construction, scoring anchors, blinding procedure, and statistical analysis. Figure 4 summarises the evaluation procedure.",
177: "For school ICT resources, availability differs from educational use, which relates to student engagement (Wang & Wang, 2023). Schools may also allocate resources in response to existing needs. Without allocation histories, the data cannot test these explanations or justify recommending less provision.",
191: "GRADE-Agent consequently offers a conversational aid for interpreting learning needs and discussing support options. Future evaluation should distinguish the quality of these replies from their contribution to educators' understanding and support decisions. This distinction would connect further assessment to educational use while retaining separate evidence requirements for each claim.",
}
REMOVE = {49, 60, 70, 74, 85, 87, 122, 124, 140, 141, 142, 143, 144, 146}
changes = []
for index, text in {**HEADINGS, **EDITS}.items():
    p = old[index]
    assert not p._p.xpath('.//m:oMath|.//w:drawing'), f'Unsafe text rewrite: {index}'
    before = p.text
    prop = deepcopy(p.runs[0]._r.rPr) if p.runs else None
    p.clear()
    r = p.add_run(text)
    if prop is not None:
        r._r.insert(0, prop)
    changes.append({'old_index': index, 'before': before, 'after': text})
for index in sorted(REMOVE):
    p = old[index]
    assert not p._p.xpath('.//m:oMath|.//w:drawing')
    changes.append({'old_index': index, 'before': p.text, 'after': '', 'reason': 'Heading consolidation or duplicate explanation; substantive detail retained elsewhere.'})
    p._p.getparent().remove(p._p)

# Place the empirical figure after the first synthesis paragraph, then explain
# its educational use. Keep original image size and caption together.
old[115]._p.addnext(old[117]._p)
old[117]._p.addnext(old[118]._p)

assert tables == [t._tbl.xml for t in d.tables]
assert math == [etree.tostring(x) for x in d._element.xpath('.//m:oMath')]
body = '\n'.join(p.text for p in d.paragraphs)
assert not re.search(r'Sections? 3\.[78]\b', body)
assert 'seed 20261027' not in body and 'seed 42' in body
numbered = [p.text for p in d.paragraphs if p.style.name.startswith('Heading') and re.match(r'^\d', p.text)]
assert numbered == list(HEADINGS.values()), numbered
assert old[39].text.startswith('Figure 1.')
assert next(i for i,p in enumerate(d.paragraphs) if p.text.startswith('Figure 1.')) < next(i for i,p in enumerate(d.paragraphs) if p.text == '3. Methods')

# Numeric results and native notation remain unedited; the revised interpretation
# paragraph 111 contained no numerical estimates in the source.
for index in [104, 106, 107, 109, 110, 112, 113, 151, 153, 154, 155, 156, 157, 158, 159, 163, 164, 165, 199, 200, 201]:
    assert old[index]._p.getparent() is not None
    assert not any(c['old_index'] == index for c in changes)
d.save(TARGET)
with ZipFile(SOURCE) as z1, ZipFile(TARGET) as z2:
    media1 = {name: hashlib.sha256(z1.read(name)).hexdigest() for name in z1.namelist() if name.startswith('word/media/')}
    media2 = {name: hashlib.sha256(z2.read(name)).hexdigest() for name in z2.namelist() if name.startswith('word/media/')}
    assert media1 == media2

record = {'source': str(SOURCE), 'target': str(TARGET), 'changes': changes,
          'headings_exactly_author_supplied': True, 'six_methods_subsections': True,
          'tables_unchanged': True, 'native_equations_unchanged': True, 'all_images_unchanged': True,
          'numerical_result_paragraphs_unchanged': True, 'bootstrap_seed': 42,
          'figure2_position': 'Within 4.4, after its first paragraph; original size retained.',
          'bootstrap_source': 'outputs/teacher_bootstrap_seed42_20261028/results_seed42.json',
          'supplementary_zip_v21_still_uses_previous_seed': True}
(QA / 'revision_audit.json').write_text(json.dumps(record, ensure_ascii=False, indent=2), encoding='utf-8')
notes = ['# v32 标题和章节结构调整', '',
         '采用作者提供的全部30个编号标题，保留原主标题、摘要、引言、关键词、数值结果、表格、全部5幅图及原生公式。', '',
         '## 章节分工', '',
         '- 第二章：保留预测贡献与调整后关联、形成性评价和情境理解的概念关系；去掉对具体应用输出和四维评分的重复展开。',
         '- 第三章：整理为6节。原预测、SHAP与DML合并在3.3，全部估计步骤和公式保留。3.4说明发现如何记录和用于应用；3.5保留案例与评分标准；3.6保留公平对照和统计分析。',
         '- 第四章：保留实证发现及供GRADE-Agent使用的分类；将负关联的详细教育解释集中到6.2。图2移到4.4第一段之后，保持原尺寸，减少分页留白。',
         '- 第五章：5.1聚焦具体教育问题，5.2和5.3保留对话和界面呈现；5.4简述评价流程并引用3.5—3.6，不重复完整评分锚点和统计方法；5.5的结果未改。',
         '- 第六章：保留教育解释和应用意义；6.4明确后续评价需区分回复质量与教育使用效果，具体设计限制仍集中在第七章。', '',
         '## Bootstrap', '',
         '方法中的随机种子已改为42，对应2026-10-28从原始评分表复算的50,000次配对案例Bootstrap。四项区间保留两位小数均与旧版相同，表格和图不需改动数值。原始评分和原种子结果均保留。',
         '旧 Supplementary_Evaluation_Materials_v21.zip 尚为原种子分析；本次不将旧材料包声称为已同步版本，正式投稿前需同步统计复现材料。', '',
         '## 术语', '',
         '|统一表述|处理|', '|---|---|',
         '|predictive contributions|预测模型贡献，不写为个体原因|',
         '|adjusted associations|DML调整后关联，不写为已确立因果效应|',
         '|learning-support summaries|当前应用的简短学习支持总结|',
         '|ordinary multi-turn LLM|两系统对照的基础模型|', '',
         '## 核查边界', '',
         '本次为章节整合，不是新增实验或消融研究。图1仍在第二章末。教师伦理与招募信息、图3旧截图与当前应用的对应，以及原有类别缺失值处理细节等待确认事项未被补造或掩盖。']
(FOLDER / 'GRADE-Agent_结构调整说明_v32.md').write_text('\n'.join(notes), encoding='utf-8')
print(TARGET)
print('Numbered headings:', len(numbered))
print('Tables, equations and all figure bytes preserved.')

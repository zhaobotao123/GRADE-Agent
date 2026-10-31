from pathlib import Path
from docx import Document
from docx.shared import Pt
base=Path('D:/Users/84130/Desktop/投稿/转投修改_20261022')
src=base/'GRADE-Agent_SEE_Teacher_Evaluation_v28.docx'
out=base/'GRADE-Agent_SEE_Teacher_Evaluation_v29.docx'
d=Document(src)
oldtext=[p.text for p in d.paragraphs]
oldxml=[p._p.xml for p in d.paragraphs]
for i,p in enumerate(d.paragraphs[:7]):
    f=p.paragraph_format
    if i==0:
        f.space_after=Pt(8)
    elif i==1:
        f.space_after=Pt(2)
    elif 2<=i<=5:
        f.line_spacing=1.35
    else:
        f.line_spacing=1.15
        f.space_before=Pt(6)
    f.keep_with_next=(i<6)
d.paragraphs[7].paragraph_format.page_break_before=True
assert [p.text for p in d.paragraphs]==oldtext
assert [p._p.xml for p in d.paragraphs[8:]]==oldxml[8:]
d.save(out)
print(out)

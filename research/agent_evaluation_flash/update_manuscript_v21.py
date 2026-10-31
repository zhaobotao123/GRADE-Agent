"""Update the existing manuscript without modifying empirical equations or sources."""
from pathlib import Path
import json,copy,re,hashlib,zipfile
from decimal import Decimal,ROUND_HALF_UP
from docx import Document
from docx.shared import Inches,Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE=Path(__file__).resolve().parents[1]
OUT=BASE/'outputs/manuscript_teacher_evaluation_v21_20261027'
SOURCE=Path('D:/Users/84130/Desktop/投稿/转投修改_20261022/GRADE-Agent_SEE_Figure5_Updated_v20.docx')
TARGET=SOURCE.parent/'GRADE-Agent_SEE_Teacher_Evaluation_v21.docx'
FIG=OUT/'figures'
doc=Document(SOURCE);old=list(doc.paragraphs)
original_equations=[p._p.xml for p in old if p._p.xpath('.//m:oMath')]
original_table=doc.tables[0]._tbl.xml

def text(i,s):
 p=old[i];p.clear();p.add_run(s);return p

changes={
3:"Using the Japanese sample from PISA 2022, we combined weighted predictive modelling, SHAP-based feature attribution, and cross-fitted DML to examine factors associated with mathematics achievement. We incorporated the resulting findings and their interpretive limits into GRADE-Agent (Grade-performance Reasoning, Analysis, Diagnosis and Explanation Agent), a conversational platform that gathers contextual information and provides a brief, provisional summary of issues that warrant attention and possible learning-support responses.",
4:"Mathematics self-efficacy showed the clearest convergence between predictive contribution and a stable positive DML-adjusted association, while disciplinary climate also showed a stable positive adjusted association. Support-related findings required more cautious interpretation. Two teachers independently rated 120 English final replies from GRADE-Agent and an ordinary multi-turn LLM across twelve synthetic cases and five runs per system. GRADE-Agent received higher mean ratings on all four dimensions. Its largest advantage concerned the appropriateness and actionability of suggestions (4.75 versus 3.61 on a five-point scale; paired difference 1.14, 95% case-cluster bootstrap CI 0.86–1.45). Agreement between teachers varied across dimensions.",
5:"These findings support the feasibility of translating achievement-related research into conversational educational support and provide preliminary evidence of higher teacher-rated reply quality than an ordinary LLM under the study conditions. They do not establish individual diagnostic accuracy or improvements in teaching decisions or student achievement. Evaluation with new cases, a larger teacher panel, and authentic educational use is needed.",
17:"We evaluate the revised conversational GRADE-Agent against an ordinary multi-turn LLM using the same language-model backbone, staged case information, and generation settings. Twelve researcher-constructed synthetic cases, each evaluated five times per system, produce 120 English final replies. Two teachers independently assess contextual relevance, consistency with case facts, appropriateness and actionability of suggestions, and clarity using anonymous, randomised materials. The comparison concerns the quality of the final replies from the complete systems, not the isolated effect of DML or an individual component. The cases had been used during earlier development and therefore do not constitute an unseen test set. The evaluation does not establish improvements in teaching decisions or student achievement.",
21:"RQ3. Compared with an ordinary multi-turn LLM receiving the same case information, how does GRADE-Agent perform in teachers' ratings of contextual relevance, consistency with case facts, appropriateness and actionability of suggestions, and clarity?",
22:"This study makes two connected contributions. First, it combines predictive modelling, SHAP-based interpretation, and DML to examine student-performance factors from complementary analytical perspectives, distinguishing predictive contributions from adjusted relationships under stated assumptions. Its contribution lies in the integration and application of these approaches, rather than in proposing a new DML estimator. Second, it develops and evaluates GRADE-Agent as a conversational platform for communicating and exploring the empirical findings. A controlled comparison with an ordinary multi-turn LLM assesses whether the resulting replies offer relevant, fact-consistent, appropriate, and understandable educational support. Together, these contributions connect student-performance analysis with an evaluated form of interactive access while preserving the distinction between population-level associations and individual educational judgement.",
38:"The value of this approach must be assessed rather than assumed. The present evaluation therefore compares the revised conversational GRADE-Agent with an ordinary multi-turn LLM under matched case-information and generation conditions. Teachers assess the final replies on separate dimensions, allowing strengths and weaknesses to be examined without assuming that a single overall score captures educational usefulness. This comparison evaluates the complete application; it cannot identify which component accounts for any observed difference.",
39:"Figure 1 summarises the three-stage research design. Stage 1 examines achievement-related factors through weighted prediction, SHAP-based attribution, and a separate DML analysis of prespecified factors in the Japanese PISA 2022 sample. Stage 2 incorporates the findings and their limits into a conversational application that gathers contextual information and offers brief, provisional learning-support summaries. Stage 3 compares its English final replies with those of an ordinary multi-turn LLM through independent teacher ratings. The final stage evaluates reply quality in synthetic cases, not changes in classroom practice or student outcomes.",
41:"This study combined an analysis of student mathematics achievement with the development and evaluation of a conversational educational-support platform. Weighted predictive modelling and SHAP-based feature attribution examined the contributions of student and contextual characteristics to achievement prediction. In parallel, DML estimated the adjusted relationships of 12 prespecified candidate factors with achievement. These findings informed GRADE-Agent's discussion of possible learning needs and support options. The revised conversational system was then compared with an ordinary multi-turn LLM using synthetic cases and independent teacher ratings. The empirical analysis and the evaluation of conversational replies address distinct questions and are reported separately.",
43:"Figure 1. Overview of the empirical analysis, conversational application, and teacher evaluation. The empirical findings concern population-level predictive patterns and adjusted associations. The conversational application relates them cautiously to user-supplied circumstances. The evaluation compares 120 English final replies from two systems across 12 synthetic cases and five runs per system, using two teachers' independent ratings.",
90:"3.7. Researcher-constructed cases and reply-quality criteria",
91:"The cases were constructed by the research team and did not use real student records. Earlier development included eight initial cases and a further set of twelve synthetic cases covering achievement changes, learning confidence, independent work, language of instruction, home resources, support timing, conflicting information, missing information, and safeguarding concerns. The present comparison reused the twelve-case set after development of the conversational prototype. These cases were known to the development process and are not described as held out.",
92:"Each case was converted into four fixed English user turns. The first three introduced the case information in stages; the fourth requested one natural paragraph of approximately 80–140 words explaining what deserved attention and what could be tried first. Both systems received the same staged facts, including safety information in the initial turn when specified. Subsequent facts were fixed rather than adapted to the questions generated by either system. The design therefore assessed the final reply after matched information exposure, not the effectiveness of adaptive information elicitation.",
93:"The rating framework drew on general distinctions between relevance, factual consistency, and understandable expression in text-generation evaluation (Fabbri et al., 2021; Liu et al., 2023), with suggestion appropriateness added for the educational task. This was an adapted task-specific rubric, not a previously validated educational diagnostic scale or an implementation of automated G-Eval scoring. Each dimension was rated separately on an integer scale from 1 to 5, with anchors at 1, 3, and 5; intermediate scores represented intermediate quality. No weighted total score was calculated.",
94:"Contextual relevance concerned whether the reply focused on issues supported by the case. A score of 1 indicated that it missed the main concern or was largely irrelevant; 3 indicated that it addressed the concern but included generic or distracting material or omitted important points; and 5 indicated a focused response that did not automatically treat neutral background information as a problem.",
95:"Consistency with case facts concerned faithful use of the supplied information. A score of 1 indicated fabrication, contradiction of key facts, or presentation of an unsupported individual cause as certain; 3 indicated broad consistency with important unsupported inferences or insufficient qualification; and 5 indicated faithful interpretation that distinguished unknown information from absent information and stayed within the support provided by the case.",
96:"Appropriateness and actionability of suggestions was the primary outcome. A score of 1 indicated unsuitable, unsafe, or unusable advice; 3 indicated generally appropriate but insufficiently specific or poorly contextualised suggestions; and 5 indicated feasible, case-sensitive suggestions that respected uncertainty. Where information was insufficient or safety was at issue, recommending further inquiry or protection could merit a high score without proposing an immediate instructional intervention.",
97:"Clarity concerned whether teachers or parents could follow the reply. A score of 1 indicated difficult-to-understand expression or organisation; 3 indicated generally understandable writing with unnecessary terminology, repetition, or organisational problems; and 5 indicated clear, coherent, and concise wording. The remaining three dimensions were secondary outcomes. Evaluation procedures and the paired analysis are detailed in Sections 5.4 and 5.5.",
98:"Two teachers rated every original final reply using independently randomised anonymous A/B materials. The materials included the complete case facts and English reply text. Chinese translations were supplied as reading aids, with instructions to judge the original English, particularly for clarity. The scores reflect these bilingual evaluation conditions rather than an independently verified English-only assessment.",
178:"Figure 1 links two distinct activities: analysing factors associated with achievement and evaluating how a conversational tool discusses learning circumstances. The empirical stage identifies predictive patterns and adjusted associations. The application stage makes these findings accessible within a dialogue. The evaluation stage asks whether teachers judge the resulting final replies more favourably than replies from an ordinary multi-turn LLM receiving the same information.",
183:"The teacher comparison provides preliminary support for the quality of the revised conversational replies under controlled conditions. Both teachers rated GRADE-Agent's suggestions more highly on average, and the paired case-level advantage was positive across all twelve cases. This supports a bounded claim about perceived appropriateness and actionability, rather than proof that the suggested actions would improve learning or that the system identifies the true cause of a student's difficulty.",
191:"6.3 What the comparison with an ordinary LLM tells us",
192:"GRADE-Agent received higher mean teacher ratings on each of the four dimensions. The clearest advantage was in the appropriateness and actionability of suggestions, which was the primary outcome. This result addresses an educationally relevant question: whether the final reply offers a plausible and usable next step in the circumstances described, rather than merely reproducing statistical terminology.",
193:"The comparison also separates score differences from agreement between teachers. Both teachers' mean ratings favoured GRADE-Agent, but weighted agreement was much higher for relevance and suggestion appropriateness than for case-fact consistency and clarity. A higher system mean does not imply that the raters applied every criterion in the same way. The latter two dimensions require more cautious interpretation and clearer calibration in future evaluations.",
194:"These findings apply to the complete conversational application. The systems differed in their task-specific prompts, use of empirical findings, and associated response-handling procedures. Because the present experiment did not remove these components separately, it cannot establish that DML information, contextual questioning, or self-checking independently caused the advantage. Earlier report-based development experiments are not combined with this comparison because the interaction design, model conditions, outputs, and evaluation criteria differ.",
195:"Providing both systems with the same staged information reduced one source of unfairness: an ordinary LLM was not restricted to an initial description while the agent received additional facts. However, the fixed information sequence also means that the study did not test whether the agent would elicit the right information from a teacher in practice. That question requires interactive user evaluation rather than scripted case completion.",
196:"The case-cluster intervals quantify variation across these twelve synthetic cases while retaining the paired comparison. They do not measure uncertainty over the wider population of teachers or establish performance on unseen cases. The two raters and the development-known cases limit the strength of any generalisation beyond this benchmark.",
197:"For educational evaluation, the useful conclusion is therefore specific: the revised research-informed conversational application produced final replies that these teachers rated more favourably than the matched ordinary LLM, especially for suggestions. This warrants further investigation of the application, without treating teacher-rated reply quality as a substitute for diagnostic validation or classroom outcomes.",
202:"The present teacher assessment concerns the relevance, factual consistency, appropriateness, and clarity of generated replies. It does not establish whether teachers would find the live interface usable, whether suggestions improve professional judgement, or whether their use changes classroom practice. The evaluation used synthetic cases and did not measure the quality of diagnosis for real students.",
203:"Further research should recruit a larger and more diverse teacher panel, use genuinely new cases, and observe teachers interacting with the system. Such studies could examine whether contextual questions are useful, whether the summaries are understandable during educational work, and whether they contribute to better-informed discussions of learning needs. Studies of actual support decisions and student outcomes would require a separate design.",
209:"Third, the conversational evaluation reused twelve researcher-constructed cases that had been available during development. Five runs per system increased the number of outputs but did not create additional independent cases. The fixed sequence of case facts did not test adaptive questioning, and the complete-system comparison did not isolate the contribution of DML or any other component. The findings should not be interpreted as performance on unseen students or evidence of component-specific effects.",
210:"Fourth, the teacher evaluation involved two raters and an adapted task-specific rubric rather than a validated diagnostic instrument. Agreement was limited for case-fact consistency and clarity, despite higher mean scores for GRADE-Agent. Chinese reading aids may also have influenced interpretation of the English replies. The bootstrap intervals are conditional on these two teachers and do not represent uncertainty over a teacher population. Future work should use additional independently recruited raters, clearer calibration, and larger, previously unseen case sets before testing the system in authentic educational settings.",
213:"Across twelve synthetic cases, five runs per system, and two independent teacher raters, GRADE-Agent received higher mean final-reply ratings than an ordinary multi-turn LLM on contextual relevance, consistency with case facts, appropriateness and actionability of suggestions, and clarity. Its largest advantage concerned suggestions (4.75 versus 3.61 out of 5; paired difference 1.14, 95% case-cluster bootstrap CI 0.86–1.45). Agreement between teachers varied across criteria. These findings support a preliminary advantage in teacher-rated reply quality under the study conditions, not individual diagnostic accuracy or improved educational outcomes.",
218:"The existing analysis software is provided as Supplementary Software 1. Supplementary Evaluation Materials provides the frozen English comparison protocol, original final replies, anonymous teacher scores, case-level source data, and the analysis and plotting script for the revised conversational evaluation. The evaluation materials are distinct from earlier report-based development records. PISA 2022 public-use microdata are not redistributed and can be obtained from the OECD PISA database.",
}
for i,s in changes.items():text(i,s)
for i in (99,100):old[i]._p.getparent().remove(old[i]._p)
text(101,'3.8. Matched system comparison and analysis')
text(102,'The formal comparison used an ordinary multi-turn LLM and the complete revised conversational GRADE-Agent, with the same DeepSeek Flash backbone and generation settings. Each system produced five conversations per case, giving 60 final replies per system and 120 in total. The frozen protocol specified suggestion appropriateness and actionability as the primary outcome; the other three rating dimensions were secondary. Earlier report-based component comparisons were not pooled with this experiment.')
text(103,'Two teachers independently rated all final replies in randomised anonymous A/B materials. Both systems received the same staged case information, and teacher scores were matched to the original replies using the concealed assignment key rather than their display labels. Implementation, language controls, and rating procedures are described in Section 5.4.')
text(104,'For each reply, the two teacher scores were averaged, followed by averaging the five runs within each case and system. Paired GRADE-Agent-minus-ordinary-LLM differences were then calculated for the twelve cases. We estimated 95% percentile bootstrap intervals by resampling the twelve paired cases with replacement 50,000 times (seed 20261027), retaining both systems and all runs within each selected case. Repeated generations were not treated as independent student cases.')
text(105,'The intervals are conditional on the two fixed teachers. Secondary-outcome intervals are exploratory, pointwise, and not adjusted for multiple comparisons. Quadratic-weighted Cohen\'s kappa and exact agreement were calculated separately for each criterion across the 120 matched replies. Mean system differences and inter-rater agreement answer different questions and are reported separately.')
text(106,'This design evaluates the quality of final conversational replies under matched information exposure. It does not isolate individual component effects, test adaptive information elicitation, or establish diagnostic accuracy or educational outcomes for real students.')

def replace_image(i,path,width):
 p=old[i];p.clear();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 p.add_run().add_picture(str(path),width=Inches(width));p.paragraph_format.keep_with_next=True
replace_image(42,FIG/'Figure1_updated_framework.png',6.20)

# Remove only the old Chapter 5 body, retaining the surrounding chapter headings.
start,end=old[124]._p,old[176]._p
node=start.getnext()
while node is not end:
 following=node.getnext();node.getparent().remove(node);node=following
anchor=old[176]

def para(s='',style='Normal',caption=False):
 p=anchor.insert_paragraph_before(s,style)
 f=p.paragraph_format
 if style=='Normal':
  p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY;f.keep_with_next=False;f.keep_together=False
  f.space_before=Pt(0);f.space_after=Pt(5);f.line_spacing=1.3
 if caption:
  f.line_spacing=1.05;f.space_after=Pt(7);p.runs[0].font.size=Pt(10)
 return p

def heading(s):return para(s,'Heading 2')
def figure(file,width,caption):
 p=para();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
 p.add_run().add_picture(str(file),width=Inches(width));para(caption,caption=True)

heading('5.1 From empirical findings to conversational learning support')
para('GRADE-Agent uses the findings in Section 4 to guide a conversation about mathematics learning rather than to assign an individual diagnosis. The empirical material distinguishes predictive contribution from DML-adjusted associations and retains uncertainty, stability, and contextual limitations. It is incorporated into the agent instructions as research-informed guidance; the comparison does not involve a separately evaluated retrieval model.')
para('Users begin by describing their observations. The agent asks focused questions about what the student can do independently, when the difficulty appeared, classroom and home circumstances, and the timing of existing support. The purpose is to clarify the situation before offering possible next steps, not to fill every research variable or rank a student using population coefficients. Neutral or protective circumstances should not automatically be labelled as causes of difficulty.')
para('When enough information is available, the agent gives a brief natural-language summary directly in the chat. It identifies what deserves attention, suggests an appropriate first step, and explains what remains uncertain. Users can continue the conversation or provide new information. The revised output is a provisional conversational reply, not a formal diagnostic report, risk score, or assertion that a statistical association explains an individual student.')

heading('5.2 Maintaining appropriate interpretations and responses')
para('The instructions require the agent to distinguish observed facts from possible explanations and to avoid treating SHAP contributions or DML estimates as individual causal findings. For example, teacher support introduced after an achievement decline cannot be described as causing that decline merely because its population-level adjusted association is negative. The reply should instead consider the chronology and what the current support involves.')
para('The revised implementation asks the model to check its wording and suggestions within the same generation call. This is a prompted self-check, not an independent critic model or a separately generated review. A local safety route can respond to explicit safety concerns in the latest user message. Structured-output validation permits one format-repair request when required; it does not provide evidence that the educational interpretation is correct.')
para('These safeguards describe intended behaviour and implementation constraints. Their individual benefits are not isolated by the present two-system evaluation. Teachers assessed the final visible replies rather than internal self-checks or a claimed independent review stage.')

heading('5.3 Presenting support through the conversational prototype')
para('The revised browser prototype presents an ongoing conversation and a short summary in the same chat area. It retains the available conversation context so that users can add observations and revise the discussion. Figure 3 retains the original prototype interface to illustrate the visual design and development history. Its report-oriented screens precede the revised conversational output evaluated here and are not screenshots of the formal English comparison.')
p=para();p._p.getparent().replace(p._p,copy.deepcopy(old[141]._p))
para('Figure 3. Original GRADE-Agent prototype interface, retained to illustrate the platform\'s visual design and earlier interaction workflow. The report-oriented screens show the earlier prototype, not the final replies assessed in this study. The revised evaluation used brief English conversational summaries and did not evaluate the interface itself.',caption=True)
para('The interface is intended to make the discussion accessible to users who do not routinely interpret statistical estimates. The teacher evaluation assessed the original final reply texts rather than live interaction with this interface. It therefore cannot establish the usability of the browser prototype or its value in an actual teacher–student consultation.')

heading('5.4 Controlled English comparison and teacher assessment')
para('The final comparison used twelve synthetic cases, two systems, and five runs per case and system, yielding 120 conversations and 60 final replies per system. Each conversation contained four fixed English user turns. Both systems used the API model alias deepseek-flash, temperature 0.25, a maximum output-token setting of 1,400, and disabled extended thinking. The model alias, prompts, case facts, and generation policy were frozen before the final run. The alias identifies the model requested from the provider; it should not be interpreted as a verified immutable model version.')
para('The ordinary LLM was a general helpful multi-turn assistant. GRADE-Agent used the revised application instructions and response handling described above. Both received identical case facts, user-turn budgets, and the final request for an 80–140-word English paragraph. A shared English-language requirement was applied to both systems, including an evaluation-only language override. The same single format-repair allowance was available to both. Case-specific expected answers or scoring constraints were not supplied to either system. Figure 4 summarises the comparison.')
figure(FIG/'Figure4_evaluation_protocol.png',6.20,'Figure 4. Protocol for the English final-reply comparison. Both systems received the same staged case information and matched generation settings. Each produced 60 final replies. Two teachers independently rated the original replies in randomised anonymous A/B materials; system identity and order were independently permuted for each teacher.')
para('An earlier English-language pilot exposed language-control and formatting problems. The common English requirement was strengthened, the revised protocol was frozen, and the entire final batch was regenerated rather than selectively replacing unsuccessful pilot outputs. The pilot was not pooled with the final comparison. All 120 final conversations completed, and the script detected no Chinese-character flags in the final batch. This automated check is not a separate certification of linguistic quality.')
para('Each teacher received all 120 original final replies arranged as 60 anonymous A/B pairs with the full case facts. Pair order and A/B assignment were independently randomised for each teacher, with each system assigned to A in 30 pairs. The original reply texts were not summarised or rewritten for rating. Chinese reading aids accompanied the English originals. The four anchored criteria described in Section 3.7 were scored separately; the instructions specified appropriateness and actionability as the primary outcome and did not request a composite total.')
para('The first teacher independently reviewed and corrected their ratings without seeing the second teacher\'s scores or knowing system identities, as confirmed by the authors. The final confirmed files were frozen before the analysis reported here. Both teachers completed all 480 required ratings, providing 960 dimension-level scores in total. No missing ratings were imputed and no teacher scores were changed by the analysis script.')
para('The analysis followed the paired case-level procedure in Section 3.8. All four rating dimensions were retained, without combining them into a weighted ranking. Results are reported as mean ratings, paired differences with case-cluster bootstrap intervals, and criterion-specific agreement between teachers.')

heading('5.5 Teacher-rated quality of the final replies')
para('GRADE-Agent received higher mean ratings on all four criteria (Table 2 and Figure 5). Its largest advantage was in the primary outcome, appropriateness and actionability of suggestions: 4.75 versus 3.61 out of 5, with a paired mean difference of 1.14 points (95% case-cluster bootstrap CI 0.86–1.45). The mean paired difference on this criterion was positive in all twelve cases. These observations concern teachers\' judgements of the replies, not measured effectiveness of the proposed actions.')
para('Table 2. Teacher ratings of English final replies and paired case-level differences.',caption=True)
table=doc.add_table(rows=1,cols=5);anchor._p.addprevious(table._tbl);table.autofit=False
widths=[2.06,.72,.72,1.66,.95]
headers=['Criterion','Ordinary\nLLM','GRADE-\nAgent','Paired difference\n[95% CI]','Weighted\nκ']
for c,s,w in zip(table.rows[0].cells,headers,widths):c.text=s;c.width=Inches(w)
summary=json.loads((OUT/'results.json').read_text(encoding='utf-8'))['summary']
def num(v,n=2):return str(Decimal(str(round(float(v),12))).quantize(Decimal('1').scaleb(-n),rounding=ROUND_HALF_UP))
for r in summary:
 cells=table.add_row().cells
 vals=[r['dimension'],num(r['ordinary_mean']),num(r['agent_mean']),f"{num(r['paired_difference'])} [{num(r['ci_low'])}, {num(r['ci_high'])}]",num(r['quadratic_weighted_kappa'],3)]
 for c,s,w in zip(cells,vals,widths):c.text=s;c.width=Inches(w)
pr=table._tbl.tblPr
existing=pr.find(qn('w:tblBorders'))
if existing is not None:pr.remove(existing)
borders=OxmlElement('w:tblBorders')
for name in ['top','bottom','left','right','insideH','insideV']:
 e=OxmlElement('w:'+name);e.set(qn('w:val'),'single' if name in ['top','bottom'] else 'nil');e.set(qn('w:sz'),'10');borders.append(e)
pr.append(borders)
for ri,row in enumerate(table.rows):
 trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
 if ri==0:trpr.append(OxmlElement('w:tblHeader'))
 for ci,cell in enumerate(row.cells):
  for p in cell.paragraphs:
   p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 else WD_ALIGN_PARAGRAPH.CENTER
   p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4);p.paragraph_format.line_spacing=1
   for run in p.runs:run.font.name='Times New Roman';run.font.size=Pt(10);run.bold=(ri==0)
  if ri==0:
   cp=cell._tc.get_or_add_tcPr();bs=OxmlElement('w:tcBorders');e=OxmlElement('w:bottom');e.set(qn('w:val'),'single');e.set(qn('w:sz'),'6');bs.append(e);cp.append(bs)
para('Note. Values are mean ratings on a 1–5 scale, averaged across two teachers and 60 replies per system (12 cases × 5 runs). Paired differences compare case-level means; positive values favour GRADE-Agent. Intervals resample 12 paired cases, with 50,000 replicates, conditional on the two teachers. Suggestion appropriateness and actionability was the primary outcome; other intervals are exploratory and unadjusted. Weighted κ is quadratic-weighted Cohen\'s kappa across the 120 replies.',caption=True)
para('The paired differences were 0.50 points for contextual relevance (95% CI 0.35–0.63), 0.19 for consistency with case facts (0.10–0.28), and 0.35 for clarity (0.24–0.48). Both teachers\' average scores favoured GRADE-Agent on each dimension. For the primary outcome, Teacher 1 rated GRADE-Agent and the ordinary LLM at 4.90 and 3.57, respectively, while Teacher 2 rated them at 4.60 and 3.65. The advantage was therefore not restricted to one teacher\'s mean ratings.')
figure(FIG/'Figure5_teacher_comparison.png',6.20,'Figure 5. Teacher-rated final-reply quality. (a) Mean ratings from two teachers for 60 original replies per system. (b) Mean paired differences after averaging teachers and five runs within each of 12 cases; error bars show 95% paired case-cluster percentile bootstrap intervals (50,000 replicates). Suggestion appropriateness and actionability is the primary outcome; the other dimensions are exploratory. The intervals are conditional on these two teachers and these development-known synthetic cases.')
para('Inter-rater agreement differed substantially by criterion. Quadratic-weighted κ was 0.794 for relevance and 0.797 for suggestion appropriateness, but 0.178 for case-fact consistency and 0.116 for clarity. Exact agreement was 88.3%, 70.8%, 74.2%, and 63.3%, respectively; all disagreements were within one scale point. The concentration of ratings near the upper scale categories can coexist with low chance-corrected agreement. These results do not justify describing agreement as uniformly strong, even though both teachers\' mean ratings favoured GRADE-Agent.')
para('The complete-system comparison supports a preliminary advantage in the perceived quality of conversational support, particularly the appropriateness of suggestions. It does not isolate an effect of DML, establish that every individual reply is superior, or validate recommendations against real student outcomes. The earlier five-configuration report benchmark is retained as a separate development record and is not used to estimate the effects reported in this section.')

# Let the interpretation paragraphs fill the preceding page instead of leaving a
# large gap when the figure-and-caption block moves to the following page.
fig5caption=next(p for p in doc.paragraphs if p.text.startswith('Figure 5. Teacher-rated'))
fig5image=fig5caption._p.getprevious()
anchor._p.addprevious(fig5image);anchor._p.addprevious(fig5caption._p)
old[206].paragraph_format.keep_with_next=True

# Add only the references used for the adapted rubric.
for index,s in [
 (229,'Fabbri, A. R., Kryściński, W., McCann, B., Xiong, C., Socher, R., & Radev, D. (2021). SummEval: Re-evaluating summarization evaluation. Transactions of the Association for Computational Linguistics, 9, 391–409. https://doi.org/10.1162/tacl_a_00373'),
 (242,'Liu, Y., Iter, D., Xu, Y., Wang, S., Xu, R., & Zhu, C. (2023). G-Eval: NLG evaluation using GPT-4 with better human alignment. In Proceedings of the 2023 Conference on Empirical Methods in Natural Language Processing (pp. 2511–2522). Association for Computational Linguistics. https://doi.org/10.18653/v1/2023.emnlp-main.153')]:
 p=old[230 if index==229 else index].insert_paragraph_before(s)
 if old[242]._p.pPr is not None:p._p.insert(0,copy.deepcopy(old[242]._p.pPr))

# Uniform changed body text, while retaining the source's heading hierarchy and equations.
old[221].paragraph_format.page_break_before=True
reference_section=False
for p in doc.paragraphs:
 if p is old[221] or p.text=='References':reference_section=True
 if reference_section and p.text!='References':p.paragraph_format.keep_together=True
 if p._p.xpath('.//m:oMath'):continue
 if p.style.name=='Normal' and p.text and not p._p.xpath('.//m:oMath'):
  p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
 for r in p.runs:
  if not r._r.xpath('.//m:oMath'):
   r.font.name='Times New Roman'
assert original_equations==[p._p.xml for p in doc.paragraphs if p._p.xpath('.//m:oMath')]
assert original_table==doc.tables[0]._tbl.xml
doc.core_properties.author='';doc.core_properties.last_modified_by='';doc.core_properties.comments=''
doc.core_properties.title=doc.paragraphs[0].text
doc.save(TARGET)
(OUT/'qa/revision_contract.json').write_text(json.dumps({'source':str(SOURCE),'output':str(TARGET),'claim':'Higher teacher-rated final-reply quality in a matched synthetic-case comparison; primary advantage in suggestion appropriateness.','not_claimed':['isolated DML/component effects','unseen-case generalisation','real-student diagnostic accuracy','educational outcome improvement'],'source_equations_preserved':True,'empirical_table_preserved':True,'n_cases':12,'n_final_replies':120,'n_teachers':2,'rating_cells':960,'pending_author_information':['Teacher recruitment, teaching background, and applicable ethics/consent arrangements require author verification before submission.'],'no_teacher_scores_changed':True},indent=2),encoding='utf-8')
print(TARGET)

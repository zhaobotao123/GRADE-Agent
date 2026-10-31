from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.oxml.ns import qn
from zipfile import ZipFile
from lxml import etree
import json,re,hashlib
from PIL import Image
import io

FOLDER=Path('D:/Users/84130/Desktop/投稿/转投修改_20261022')
SOURCE=FOLDER/'GRADE-Agent_SEE_NatureSkills_Polished_v30.docx'
TARGET=FOLDER/'GRADE-Agent_SEE_Refined_v31.docx'
QA=Path('outputs/manuscript_revision_v31_20261027')
audit=json.loads((QA/'reference_audit.json').read_text(encoding='utf-8'))
d=Document(SOURCE)
old=list(d.paragraphs)
table_xml=[t._tbl.xml for t in d.tables]
math_xml=[etree.tostring(x) for x in d._element.xpath('.//m:oMath')]
edits={
24:'2.1 Prediction, feature attribution, and adjusted-association estimation',
67:'We compared Random Forest, ExtraTrees, RidgeCV, ElasticNetCV, and LightGBM (Ke et al., 2017) using the same 42-variable feature set and training–test split for each plausible value. Five-fold cross-validation selected LightGBM hyperparameters.',
99:'Both systems received identical staged information. Two teachers independently rated the original English final replies alongside full case descriptions in anonymous A/B pairs. Pair order and system assignment were randomised separately for each teacher. A concealed assignment key matched ratings to source replies rather than display labels.',
119:'Figure 2. Integration of predictive, adjusted-association, and rule-assignment evidence for GRADE-Agent. (a) Weighted performance of the five candidate models. (b) Predictive contributors ranked by mean absolute SHAP value. (c) Pooled DML estimates with 95% confidence intervals. (d) Integration of predictive and adjusted-association evidence into four diagnostic rule classes. SHAP describes model use, not causal interpretation.',
138:"Teacher evaluation examined whether the final replies helped educators interpret a learning situation and consider appropriate support. The comparison concerned educational reply quality, beyond the presence of statistical information.",
139:"The evaluation compared 120 English final replies across 12 synthetic cases and five runs per system. Both systems received the same case information. Teachers judged the replies in relation to the full case descriptions.",
146:'All dimensions used a five-point scale, with suggestion appropriateness and actionability as the prespecified primary outcome. The remaining dimensions addressed complementary aspects of educational reply quality.',
147:'Sections 3.7 and 3.8 provide the case construction, scoring anchors, blinding procedure, and statistical analysis. Here, the four dimensions organise the interpretation of teacher ratings in relation to educational support.',
188:'6.4 Implications for educational evaluation and learning support',
189:'The findings suggest a role for achievement analysis in educational evaluation that extends beyond describing performance differences. Predictive patterns and adjusted associations can guide questions about learning circumstances, helping educators decide what information deserves closer attention.',
190:'For learning support, the practical step is to connect a research finding with observations that educators can examine. Questions about independent work, responses to assistance, and the timing of difficulties make the discussion relevant to support planning. They also leave room for alternative explanations as information develops.',
191:"Professional judgement gives these observations their educational meaning. Teachers can relate a reply to curriculum demands, classroom relationships, health, and existing support arrangements that the statistical data do not capture. Research-informed conversation should help organise that judgement while preserving the teacher's responsibility for action.",
192:'The four evaluation dimensions also offer a way to examine the quality of AI-supported educational discussion. Relevance concerns attention to the learning situation, factual consistency concerns faithful interpretation, actionability concerns possible next steps, and clarity concerns understandable communication. Considering them separately avoids equating fluent expression with educational usefulness.',
194:'GRADE-Agent consequently offers a conversational aid for interpreting learning needs and discussing support options. Its role is to make analytical findings usable alongside professional knowledge of the learner, with recommendations remaining open to review.',
198:"Second, the Japanese sample defines the initial educational context. School organisation, teaching practices, resources, and self-report patterns may affect associations in other systems. Some self-reported measures may also incompletely represent learning circumstances. Cross-system replication should reanalyse relevant data and reassess the evidence base rather than assume unchanged relationships.",
199:'Third, the evaluation reused twelve researcher-constructed cases available during development. Five runs per system increased outputs, not independent cases. Fixed facts did not test adaptive questioning, and the complete-system comparison did not isolate DML or other components. The ratings concern generated replies, not live-interface usability, classroom practice, or diagnostic accuracy for real students. Future work should use unseen cases and observe teachers interacting with the system to examine contextual questions and support planning. Effects on decisions and student outcomes require a separate study design.'
}
edits[200]='Fourth, two teachers used an adapted task-specific rubric, not a validated diagnostic instrument. Agreement was limited for factual consistency and clarity, despite higher mean ratings for GRADE-Agent. Chinese reading aids may have influenced interpretation of English replies. Bootstrap intervals are conditional on these teachers, without representing uncertainty across a teacher population. Future studies should recruit a larger, more diverse teacher panel, clarify scoring anchors, and evaluate previously unseen cases before authentic classroom use.'
removed_indices={96,140,193}
changes=[]
for i,t in edits.items():
    p=old[i];before=p.text
    prop=deepcopy(p.runs[0]._r.rPr) if p.runs else None
    p.clear();r=p.add_run(t)
    if prop is not None:r._r.insert(0,prop)
    changes.append({'paragraph':i,'before':before,'after':t})
for i in removed_indices:
    p=old[i]
    changes.append({'paragraph':i,'before':p.text,'after':'','reason':'Repetition removed; relevant conditions retained in Methods or Limitations.'})
    p._p.getparent().remove(p._p)

# Remove bibliography entries only after checking main and available supplement narratives.
removed_refs=[]
for ref in audit['references']:
    if not ref['body_match'] and not ref['supplement_matches'] and ref['first_author']!='Ke':
        p=old[ref['paragraph']]
        removed_refs.append(ref)
        p._p.getparent().remove(p._p)
    elif ref['first_author']=='Ke':
        p=old[ref['paragraph']]
        r=p.add_run(' https://proceedings.neurips.cc/paper_files/paper/2017/hash/6449f44a102fde848669bdd9eb6b76fa-Abstract.html')
        if p.runs[0]._r.rPr is not None:r._r.insert(0,deepcopy(p.runs[0]._r.rPr))

assert len(removed_refs)==14
assert table_xml==[t._tbl.xml for t in d.tables]
assert math_xml==[etree.tostring(x) for x in d._element.xpath('.//m:oMath')]
assert not any('adjusted-effect' in p.text.lower() for p in d.paragraphs)
# Replace Figure 2 raster and SVG coherently, preserving width and aspect ratio.
shape=d.inline_shapes[1]
rid=shape._inline.graphic.graphicData.pic.blipFill.blip.embed
part=d.part.related_parts[rid]
png=(FOLDER/'Figures_v31/Figure2_adjusted_associations.png').read_bytes()
part._blob=png
im=Image.open(io.BytesIO(png))
height=int(shape.width*im.height/im.width)
shape.height=height
for extent in shape._inline.xpath('.//a:xfrm/a:ext'):
    extent.set('cy',str(height))
# Word may use an SVG extension instead of the raster fallback.
for elem in shape._inline.iter():
    if etree.QName(elem).localname=='svgBlip':
        srid=elem.get(qn('r:embed'))
        d.part.related_parts[srid]._blob=(FOLDER/'Figures_v31/Figure2_adjusted_associations.svg').read_bytes()
d.save(TARGET)

nd=Document(TARGET)
refs_start=next(i for i,p in enumerate(nd.paragraphs) if p.text=='References')
refs=[p.text for p in nd.paragraphs[refs_start+1:] if p.text]
assert len(refs)==26
body='\n'.join(p.text for p in nd.paragraphs[:refs_start])
assert 'Ke et al., 2017' in body
assert 'Chinese reading aids' in body
assert not any('Chinese' in p.text for p in nd.paragraphs if p.text in [edits[138],edits[139],edits[146],edits[147]])
record={'source':str(SOURCE),'target':str(TARGET),'changes':changes,'references_before':40,'references_after':26,'removed_uncited_references':removed_refs,'ke_reference_retained_and_cited':True,'tables_unchanged':True,'equations_unchanged':True,'figure2_data_unchanged':True,'bilingual_condition_retained_in_limitations':True}
(QA/'revision_audit.json').write_text(json.dumps(record,ensure_ascii=False,indent=2),encoding='utf-8')

# A readable, itemised report preserves every removed entry and explains audit limits.
report=['# 论文修订与引文核对说明 v31','',
'## 已完成的修改','',
'1. 将正文标题及图2图注中的 adjusted-effect 统一为 adjusted-association；图2内部两处标注同步修改。原绘图数据及估计方法未变，仅对标题间距和示意框文字作适应性排版。',
'2. 3.7删除末段重复的教师材料说明；3.8保留独立评分、匿名随机呈现及统计分析。5.4仅保留评价目的、必要样本信息、四个教育维度与方法交叉引用，不再重复评分随机化和自助法细节。',
'3. 6.4集中讨论教育评价、学习支持和教师判断的启示。日本样本适用性、案例限制、系统使用评价及后续设计集中在第七章。',
'4. 中文阅读辅助不再出现在3.7和5.4，保留在第七章作为实际评价条件的简短说明。没有把真实的双语辅助评价改写成纯英文评价。','',
'## 引文核对结果','',
'核对主文正文、图注和表格，并检查可找到的补充说明文档和两个提交材料包中的叙述性README等文件。没有把原始大模型回复里的文献提及当作作者正式引用。旧版补充说明属于历史实验材料，本次没有改写它或将其混入新实验。',
'原参考文献40条，其中15条未在主文找到对应引用；所查补充叙述中也未找到这些条目的正式作者年份引用。Ke等（2017）是实际使用的LightGBM方法来源，因此补到3.3并保留；其余14条从本次主文文末移除，完整条目留在下表以备恢复。修订后主文参考文献为26条。','',
'这属于作者—年份对应与书目信息核对，不等同于逐篇全文验证所有引文是否充分支持具体论断，也不覆盖未提供的外部材料。','',
'### 从主文参考文献表移除的条目','',
'|作者年份|处理理由|完整条目|','|---|---|---|']
for ref in removed_refs:
    report.append(f"|{ref['first_author']} ({ref['year']})|所查主文和补充叙述未引；保留审计记录|{ref['entry']}|")
report+=['','### 全部40条的核对记录','','|作者年份|主文引用|所查补充引用|元数据检查|处理|','|---|---|---|---|---|']
fallback={
'Albreiki':'出版方记录可定位；CrossRef限流，见DOI原链接',
'Bond':'Crossmark出版记录与作者年份相符：https://crossmark.crossref.org/dialog/?doi=10.1186%2Fs41239-023-00436-z',
'Jin':'作者机构库记录相符：https://repository.eduhk.hk/en/publications/investigating-behavioral-patterns-to-facilitate-performance-predi/',
'Klug':'作者机构库记录相符：https://ucrisportal.univie.ac.at/en/publications/diagnostic-competence-of-teachers-a-process-model-that-accounts-f/',
'Ke':'NeurIPS正式记录相符：https://proceedings.neurips.cc/paper_files/paper/2017/hash/6449f44a102fde848669bdd9eb6b76fa-Abstract.html',
'Lundberg':'NeurIPS正式记录相符：https://proceedings.neurips.cc/paper_files/paper/2017/hash/8a20a8621978632d76c43dfd28b67767-Abstract.html',
'Hu':'in press信息由作者提供，尚无DOI，需刊出后更新',
'Cai':'未进行全文核验；主文未使用',
}
for ref in audit['references']:
    status=ref['external_metadata']['status']
    desc='CrossRef题名/首作者/年份对应' if status=='verified_title_author_year' else fallback.get(ref['first_author'],'未完全核验；CrossRef不收录或需人工确认')
    decision='保留并补正文方法引用' if ref['first_author']=='Ke' else ('保留' if ref['body_match'] or ref['supplement_matches'] else '移至本审计记录')
    report.append(f"|{ref['first_author']} ({ref['year']})|{'有' if ref['body_match'] else '未检出'}|{'有' if ref['supplement_matches'] else '未检出'}|{desc}|{decision}|")
report+=['','CrossRef有31条题名、首作者与发表年份对应成功。4条遇到限流，使用出版记录或作者机构库辅助核对；另有会议论文、预印本与作者确认的in-press条目按其实际来源区别记录。CrossRef不收录arXiv DOI不是文献不存在的证据。','',
'## 保留的待确认事项','',
'教师伦理／知情同意安排、图3旧界面与当前系统的对应、方法软件版本及类别缺失值规则，仍需依据真实材料确认。本次不补造声明，不改实验数据。']
(FOLDER/'GRADE-Agent_修订与引文核对_v31.md').write_text('\n'.join(report),encoding='utf-8')
print(TARGET)
print('References: 40 ->',len(refs),'; uncited removed:',len(removed_refs))

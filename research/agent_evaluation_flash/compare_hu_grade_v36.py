from pathlib import Path
import json,sys
from pypdf import PdfReader
from docx import Document
BASE=Path(__file__).resolve().parents[1]
OUT=BASE/'outputs/hu_grade_comparison_20261028';OUT.mkdir(parents=True,exist_ok=True)
pdf=Path('D:/微信/xwechat_files/wxid_lhoxzs8swzbs22_7edd/msg/file/2026-10/2026 Studies in Educational Evaluation.pdf')
r=PdfReader(pdf)
pages=[p.extract_text()for p in r.pages]
(OUT/'published_fulltext.txt').write_text('\n\n'.join(f'=== PDF PAGE {i+1} ===\n{t}'for i,t in enumerate(pages)),encoding='utf-8')
d=Document('D:/Users/84130/Desktop/投稿/转投修改_20261022/GRADE-Agent_SEE_Checked_v36.docx')
(OUT/'grade_fulltext.txt').write_text('\n'.join(p.text for p in d.paragraphs)+'\n'+'\n'.join(' | '.join(c.text for c in row.cells)for t in d.tables for row in t.rows),encoding='utf-8')
print('PAGES',len(pages))
start,end=(int(sys.argv[1]),int(sys.argv[2]))if len(sys.argv)>2 else(1,len(pages))
for i,t in enumerate(pages):
 if start<=i+1<=end:print(f'=== PDF PAGE {i+1} ===\n{t}')

from pathlib import Path
from copy import deepcopy
import json
from docx import Document

base=Path('D:/Users/84130/Desktop/投稿/转投修改_20261022')
source=base/'GRADE-Agent_SEE_Teacher_Evaluation_v23.docx'
target=base/'GRADE-Agent_SEE_Teacher_Evaluation_v24.docx'
d=Document(source);ps=list(d.paragraphs)
start=next(p for p in ps if p.text.startswith('6.3 What'))
end=next(p for p in ps if p.text.startswith('6.4 Implications'))
a,b=ps.index(start),ps.index(end)
before=[p._p.xml for p in ps[:a]];after=[p._p.xml for p in ps[b:]]
media=[s._inline.xml for s in d.inline_shapes]
tables=[t._tbl.xml for t in d.tables]
properties=deepcopy(ps[a+1]._p.pPr)
paragraphs=[
"The comparison with an ordinary multi-turn LLM highlights an important challenge in educational support: identifying factors associated with student achievement is only one part of the task. These findings also need to be presented in ways that help educators understand a student's circumstances and consider appropriate forms of learning support. In this study, GRADE-Agent received higher mean teacher ratings across all four evaluation dimensions, with the clearest advantage in the appropriateness and actionability of suggestions. This pattern suggests that connecting empirical findings with information about a learner's situation may help move the response beyond general advice towards more context-sensitive educational recommendations.",
"The difference was particularly evident in the primary outcome. Teachers rated GRADE-Agent's suggestions more highly in terms of their suitability for the situation described and their usefulness in identifying possible next steps. The higher ratings for contextual relevance also suggest that the responses were better aligned with the information provided in each case. Together, these findings point to the importance of connecting research on achievement-related factors with the circumstances in which learning difficulties arise. For educators, the value of such a system lies not simply in presenting research findings, but in helping relate those findings to questions that can inform further discussion and support planning.",
"The results also show that the quality of educational support has several dimensions. A recommendation may be relevant to a student's situation yet still require clearer explanation or more careful attention to the case facts. Although GRADE-Agent received higher mean ratings for factual consistency and clarity, agreement between the two teachers was lower on these dimensions than on contextual relevance and suggestion appropriateness. This distinction is important because useful educational support depends both on the relevance of a recommendation and on whether educators can clearly understand the reasoning behind it. Future evaluations would benefit from more explicit scoring guidance for factual consistency and clarity.",
"Overall, the findings suggest that bringing empirical research into an educational conversation may be most useful when research findings are connected to the learner's described circumstances and translated into practical options for support. GRADE-Agent offers one approach to making that connection. Further evaluation with teachers working through authentic educational situations would help establish how such systems can contribute to the interpretation of learning needs and the planning of appropriate support.",
]
start.clear();start.add_run('6.3. What the Comparison with an Ordinary LLM Tells Us')
node=start._p.getnext()
while node is not end._p:
 nxt=node.getnext();node.getparent().remove(node);node=nxt
for s in paragraphs:
 p=end.insert_paragraph_before(s)
 if properties is not None:
  if p._p.pPr is not None:p._p.remove(p._p.pPr)
  p._p.insert(0,deepcopy(properties))
new=list(d.paragraphs)
assert before==[p._p.xml for p in new[:a]]
newend=next(p for p in new if p.text.startswith('6.4 Implications'))
assert after==[p._p.xml for p in new[new.index(newend):]]
assert [p.text for p in new[a+1:new.index(newend)]]==paragraphs
assert media==[s._inline.xml for s in d.inline_shapes]
assert tables==[t._tbl.xml for t in d.tables]
d.save(target)
qa=Path(__file__).resolve().parents[1]/'outputs/manuscript_section63_v24_20261027/qa'
qa.mkdir(parents=True,exist_ok=True)
(qa/'verification.json').write_text(json.dumps({'source':str(source),'output':str(target),'author_section63_verbatim':True,'other_paragraph_xml_unchanged':True,'all_tables_figures_unchanged':True},ensure_ascii=False,indent=2),encoding='utf-8')
print(target)

"""Apply the author's scoped replacement instructions to preserved v37."""
from pathlib import Path
from copy import deepcopy
import re,json,zipfile,hashlib,sys,subprocess,csv
from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement
from lxml import etree
BASE=Path(__file__).resolve().parents[1]
F=Path('D:/Users/84130/Desktop/投稿/转投修改_20261022')
OUT=BASE/'outputs/author_instructions_v38_20261028';OUT.mkdir(parents=True,exist_ok=True)
REQ=Path('C:/Users/84130/.codex/attachments/516a5c84-a4c9-454a-b0c3-ac805aa420e2/Pasted text.txt').read_text(encoding='utf-8-sig')
blocks=[b.strip() for b in re.split(r'\n\s*\n',REQ)]
def supplied(prefix):
 m=[b for b in blocks if b.startswith(prefix)]
 assert len(m)==1,(prefix,len(m))
 return m[0]
source=F/'GRADE-Agent_SEE_Positioning_Revised_v37.docx';d=Document(source)
orig=list(d.paragraphs);tables=[t._tbl.xml for t in d.tables]
refs=[p._p.xml for p in orig[195:]]
math=[etree.tostring(e) for e in d._element.xpath('.//m:oMath')]
changes=[]
def put(p,text):
 old=p.text
 assert not p._p.xpath('.//w:drawing|.//m:oMath')
 prop=deepcopy(p.runs[0]._r.rPr) if p.runs else None
 p.clear();r=p.add_run(text)
 if prop is not None:r._r.insert(0,prop)
 changes.append({'before':old,'after':text})
def group(indices,texts):
 ps=[orig[i] for i in indices]
 template=deepcopy(ps[0]._p.pPr)
 runprop=deepcopy(ps[0].runs[0]._r.rPr) if ps[0].runs else None
 last=ps[0];put(last,texts[0])
 for text in texts[1:]:
  e=OxmlElement('w:p')
  if template is not None:e.append(deepcopy(template))
  last._p.addnext(e);p=Paragraph(e,last._parent);r=p.add_run(text)
  if runprop is not None:r._r.insert(0,deepcopy(runprop))
  changes.append({'before':'[new paragraph]','after':text});last=p
 for p in ps[1:]:
  changes.append({'before':p.text,'after':'[replaced/deleted]'});p._p.getparent().remove(p._p)

# Abstract: replace the final interpretation, leaving all preceding data intact.
old_end='The findings support further evaluation of research-informed conversational support, without establishing individual diagnostic accuracy or improved educational outcomes.'
assert orig[2].text.endswith(old_end)
put(orig[2],orig[2].text[:-len(old_end)]+supplied('This study demonstrates how empirical analysis'))
put(orig[5],supplied('Student achievement is an important indicator'))
put(orig[14],orig[14].text.split(' This complete-system comparison evaluates')[0])
put(orig[16],supplied('**RQ1:**').replace('**RQ1:**','RQ1:'))
group([19],[supplied('First, this study combines weighted'),supplied('Second, this study translates'),supplied('By connecting empirical analysis with contextual')])

# Full replacement of 2.2, preserving citations in the nearest related claim.
put(orig[25],'2.2 Formative Assessment and Individualized Learning Support')
p22=[supplied(s) for s in ['Formative assessment emphasizes','Similar achievement levels','Building on this perspective','The agent provides a structured reference']]
p22[-1]+=' (Bond et al., 2024; Khosravi et al., 2022; Zawacki-Richter et al., 2019).'
p22[-1]=p22[-1].replace('process. (Bond','process (Bond')
group([26,27,28,29],p22)

# Check original DML implementation and the actual selected-variable record.
codefolder=Path('C:/Users/84130/Documents/Codex/2026-07-03/zhe/pisa_grade_agent_formal_experiment')
code=(codefolder/'run_formal_pisa_experiment.py').read_text(encoding='utf-8-sig')
config=(codefolder/'config.py').read_text(encoding='utf-8-sig')
utils=(codefolder/'utils_pisa.py').read_text(encoding='utf-8-sig')
with Path('D:/pisa2022/outputs_grade_agent_formal/selected_variables.csv').open(encoding='utf-8-sig',newline='') as f:selected=list(csv.DictReader(f))
pool=[r['variable'] for r in selected if r['role'] in ['treatment','control_or_covariate']]
assert len(pool)==21 and 'CNT' in pool and 'x_cols = [c for c in feature_cols if c != treatment]' in code
assert 'OneHotEncoder(handle_unknown="ignore", sparse_output=False, min_frequency=20)' in utils
assert '"CNT"' in config and 'ID_VARS = ["CNTSTUID", "CNTSCHID"]' in config
covtext=supplied('For each DML analysis, one candidate factor')
covtext+=' The background covariates were CNT, ST004D01T, AGE, GRADE, LANGN, HISCED, PAREDINT, HISEI, and RELATST. The implemented pool contained 21 variables, with 20 non-focal variables per analysis; CNT was retained but constant within Japan.'
put(orig[55],covtext)
group([80,81],[supplied('GRADE-Agent organizes the empirical findings'),supplied('This process connects population-level evidence')])
put(orig[83],orig[83].text.split(' The present comparison reused')[0])
group([84],[supplied('Each case consisted of four fixed user turns.'),supplied('This design ensured that the two systems')])
rub=supplied('The evaluation rubric was developed around four dimensions')
rub+=' The dimensions drew on relevance, factual consistency, and understandable expression in text-generation evaluation, with recommendation appropriateness adapted to the educational task (Fabbri et al., 2021; Liu et al., 2023).'
put(orig[85],rub)
stats=[supplied(s) for s in ['This study compared an ordinary LLM','For each of the 12 synthetic cases','The appropriateness and actionability of recommendations was prespecified','To examine the stability of the observed differences','Quadratic-weighted Cohen\'s kappa was also calculated','Together, these procedures evaluate']]
stats[1]+=' Teachers viewed the original English replies alongside full case descriptions in anonymous A/B pairs. Pair order and system assignment were randomized separately for each teacher, and a concealed key matched ratings to source replies.'
stats[3]=stats[3].replace('50,000 bootstrap samples were drawn','50,000 bootstrap samples (random seed 42) were drawn')
stats[4]+=' Exact agreement was also calculated for each dimension across 120 matched replies. Confidence intervals were conditional on the two teachers; complementary-outcome intervals were exploratory, pointwise, and unadjusted for multiple comparisons.'
group([91,92,93,94,95],stats)

# Results: preserve Table2, its caption/note and Figure5 objects, replace prose only.
put(orig[137],supplied('Across all four evaluation dimensions, GRADE-Agent received'))
group([140,141,142,143,144,145],[supplied(s) for s in ['For contextual relevance, GRADE-Agent received a mean','The teacher-specific results showed','Inter-rater agreement was relatively high','Overall, the mean ratings across']])
for i in [150,151]:orig[i]._p.getparent().remove(orig[i]._p)
put(orig[159],supplied('Both teachers gave higher ratings'))
group([168,169,170,171],[supplied(s) for s in ['The comparison between GRADE-Agent and the ordinary','These results align with the design','For contextual relevance, the higher mean','Overall, the teacher evaluations provide direct']])
group([173,174,175,176,177],[supplied(s) for s in ['The findings indicate that analysis','Building on this evidence, GRADE-Agent','The teacher evaluations further showed','By connecting achievement-related analysis']])
group([180,181,182,183],[supplied(s) for s in ['First, the PISA 2022 data','Second, this study focused on students in Japan.','Third, the system evaluation used 12','Fourth, two teachers independently evaluated']])
group([185,186,187],[supplied(s) for s in ['This study examined factors associated with mathematics','The empirical analyses showed that mathematics','In the system evaluation, GRADE-Agent','By integrating achievement-related analysis']])
orig[185].paragraph_format.keep_together=True
orig[191].paragraph_format.keep_together=True

# Consistency edits authorised by the instruction; references and tables untouched.
for p in d.paragraphs:
 if p.text=='References':break
 if p._p.xpath('.//w:drawing|.//m:oMath'):continue
 t=p.text.replace('ordinary multi-turn LLM','ordinary LLM').replace('ordinary large language model (LLM)','ordinary LLM (large language model)')
 t=t.replace('Suggestion appropriateness and actionability','Appropriateness and actionability of recommendations')
 t=t.replace('suggestion appropriateness and actionability','appropriateness and actionability of recommendations')
 t=t.replace('suggestion appropriateness','recommendation appropriateness')
 t=t.replace('case-fact consistency','consistency with case facts')
 if t!=p.text:put(p,t)

assert tables==[t._tbl.xml for t in d.tables]
assert math==[etree.tostring(e) for e in d._element.xpath('.//m:oMath')]
ri=next(i for i,p in enumerate(d.paragraphs) if p.text=='References')
assert refs==[p._p.xml for p in d.paragraphs[ri+1:]]
target=F/'GRADE-Agent_SEE_Educational_Contribution_v38.docx';d.save(target)
with zipfile.ZipFile(source) as a,zipfile.ZipFile(target) as b:
 media=[n for n in a.namelist() if n.startswith('word/media/')]
 assert all(a.read(n)==b.read(n) for n in media)
full='\n'.join(p.text for p in d.paragraphs)
assert 'ordinary multi-turn LLM' not in full
assert 'random seed 42' in full
assert len(d.tables)==2
audit={'source':str(source),'output':str(target),'changes':changes,'table_xml_unchanged':True,'reference_xml_unchanged':True,'math_xml_unchanged':True,'media_unchanged':True,'DML_pool':pool,'DML_nonfocal_count':20,'DML_CNT_retained_constant':True,'code_sha256':{p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in [codefolder/'config.py',codefolder/'utils_pisa.py',codefolder/'run_formal_pisa_experiment.py']},'abstract_words':len(d.paragraphs[2].text.split())}
(OUT/'revision_audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf-8')
log='''# v38 修订记录

基于 v37，按作者提供的分节指令修改；原版保留。

## 已完成

- 替换摘要结尾、引言首段、RQ1、两项贡献，以及2.2、3.4、3.5指定段落、3.6、5.5周边文字、6.1末段、6.3、6.4、局限性及结论。
- 统一正文评价维度名称；基线称 ordinary LLM。未重写未指定的实证结果或2.1/2.3。
- 保留原有表格XML、原生公式、全部嵌入图像和参考文献条目。保留评分、区间、κ、模型统计量及样本数字；没有重跑生成或统计实验。
- 3.6保留配对案例层级、50,000次Bootstrap、随机种子42、两位教师、原始英文回复、随机匿名呈现、隐藏映射、κ及精确一致率。保留区间条件及多重比较说明，以满足“保留统计实现细节”的要求。
- 2.2与评分标准替换涉及的原引用，移至同主题新段落，避免未经要求删除引用或留下对应条目无正文引用。

## DML代码核查

原始目录：C:/Users/84130/Documents/Codex/2026-07-03/zhe/pisa_grade_agent_formal_experiment/

核查config.py、run_formal_pisa_experiment.py、utils_pisa.py，并与D:/pisa2022/outputs_grade_agent_formal/selected_variables.csv对照。已确认12个保留因素加9个背景变量组成21变量池；每次移除当前焦点因素，使用20个非焦点变量。背景变量为CNT、ST004D01T、AGE、GRADE、LANGN、HISCED、PAREDINT、HISEI、RELATST。CNT在日本样本内为常量，保留于分类预处理；不存在主动删除常量列的步骤。不能称其提供跨国调整。本次未改变代码、变量列表或结果。

## 仍需作者决定或核实

- 按要求保留表2和图形：其中原有“suggestions”或简写标签未改为正文的“recommendations”；指同一评价维度，严格字面统一需另行授权改图表。
- 按要求删去案例在开发阶段已使用的重复说明，并同步清理引言中的对应句子；未改称独立留出集。该设计事实未改变，原始方案及补充记录应准确保留。
- 当前评价为固定分阶段信息下的完整系统最终回复比较，不是组件因果贡献检验；新正文没有声称进行了消融或证实真实学生成绩提升。
- 匿名评分不等同于伦理审批/豁免或知情同意。相关信息及作者、资助、利益冲突仍待实际确认。
- 摘要按作者指定内容扩展，当前词数见revision_audit.json；期刊最新字数限制尚未独立确认，应在最终投稿前对照官方指南。
'''
(F/'GRADE-Agent_v38_修订与核查记录.md').write_text(log,encoding='utf-8')
print('Saved',target,'Abstract words',audit['abstract_words'],flush=True)
renderer=Path('C:/Users/84130/.codex/plugins/cache/openai-primary-runtime/documents/26.1025.11240/skills/documents/render_docx.py')
if '--no-render' not in sys.argv:
 subprocess.run([sys.executable,str(renderer),str(target),'--output_dir',str(OUT/'render'),'--emit_pdf'],check=True)

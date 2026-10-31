from copy import deepcopy
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

ROOT = Path('D:/Users/84130/Desktop/投稿/转投修改_20261022')
SRC = ROOT / 'GRADE-Agent_SEE_Teacher_Evaluation_v27.docx'
OUT = ROOT / 'GRADE-Agent_SEE_Teacher_Evaluation_v28.docx'
BILINGUAL = ROOT / 'GRADE-Agent_Chapter2_中英对照_v28.docx'
QA = Path('outputs/manuscript_chapter2_v28_20261027')
QA.mkdir(parents=True, exist_ok=True)
doc = Document(SRC)
paras = doc.paragraphs
start = next(i for i,p in enumerate(paras) if p.text == '2. Conceptual Framework')
end = next(i for i,p in enumerate(paras) if p.text == '3. Materials and Methods')
outside = [p._p.xml for p in paras[:start]] + [p._p.xml for p in paras[end:]]
tables = [t._tbl.xml for t in doc.tables]
images = [s._inline.xml for s in doc.inline_shapes]
old_words = sum(len(p.text.split()) for p in paras[start:end])
sections = [
    ('2.1 Prediction, feature attribution, and adjusted-effect estimation', '2.1 预测、特征归因与调整后效应估计', [
        ('The analytical framework distinguishes predictive contribution from an adjusted relationship with achievement. Predictive modelling examines how student and contextual characteristics jointly account for variation in achievement (Albreiki et al., 2021; Alnasyan et al., 2024). SHAP describes how a fitted model uses these characteristics in its predictions (Lundberg & Lee, 2017). DML addresses a separate question: how a prespecified factor relates to achievement after adjustment for the observed covariates (Chernozhukov et al., 2018).',
         '本研究的分析框架区分了预测贡献与因素同学业成绩之间的调整后关系。预测建模考察学生特征和情境特征如何共同解释学业成绩的差异（Albreiki et al., 2021; Alnasyan et al., 2024）。SHAP说明拟合后的模型如何利用这些特征作出预测（Lundberg & Lee, 2017）。DML回答另一个问题：在调整已观测协变量后，预先指定的因素与学业成绩之间存在怎样的关系（Chernozhukov et al., 2018）。'),
        ('These perspectives inform different educational interpretations. A strong predictive contribution indicates that a characteristic helps the model distinguish achievement patterns. An adjusted relationship provides an additional basis for discussing that factor under the specified assumptions. The two findings need not coincide, and neither alone establishes the benefit of changing the factor through an educational intervention.',
         '这些分析视角对应不同的教育解释。较强的预测贡献意味着某项特征有助于模型区分学业成绩模式；调整后关系则为在既定假设下讨论该因素提供了另一种依据。两类发现未必一致，也都不能单独证明通过教育干预改变该因素就会带来收益。'),
        ('This distinction guides how the findings enter educational discussion. Predictive rankings describe information used by the model, while DML estimates describe adjusted relationships for prespecified factors. The DML factors are specified independently of the SHAP ranking. Their interpretation retains the uncertainty and population-level scope of the underlying analysis.',
         '这一区分决定了研究发现如何进入教育讨论。预测排序描述模型利用的信息，DML估计则描述预先指定因素的调整后关系。DML分析因素的设定独立于SHAP排序；解释这些因素时，应保留原始分析所具有的不确定性和群体层面的适用范围。')
    ]),
    ('2.2 Formative assessment and human-centred diagnostic support', '2.2 形成性评价与以人为中心的诊断支持', []),
    ('2.3 From empirical analysis to interactive diagnostic support', '2.3 从实证分析到交互式诊断支持', [
        ('The framework connects achievement-related findings with educational support through contextual inquiry. Empirical analysis identifies factors that warrant consideration, while information about a learner helps establish which questions are relevant to the situation. For example, an association involving teacher support prompts questions about when support began and which difficulties preceded it. The association alone cannot explain the role of support in that learner\'s performance.',
         '本研究框架通过情境了解，将学业成绩相关发现与教育支持联系起来。实证分析识别值得考虑的因素，有关学习者的信息则帮助判断哪些问题与当前情境相关。例如，涉及教师支持的统计关系，会提示我们进一步了解支持何时开始，以及此前已经出现哪些学习困难。仅凭该统计关系，无法解释支持在该学习者成绩表现中所起的作用。'),
        ('Conversation provides a way to develop this contextual understanding through questions, clarification, and revision of initial interpretations (Chiu et al., 2023; Kuhail et al., 2023; Yan et al., 2024). Within GRADE-Agent, research findings inform what deserves attention, and the information supplied by users informs the discussion of possible support. The resulting summary links the available facts with issues for further consideration and possible next steps. Teachers retain responsibility for checking these interpretations and deciding whether the suggested actions suit the learner.',
         '对话通过提问、澄清和修正初步解释，为形成这种情境理解提供了一种方式（Chiu et al., 2023; Kuhail et al., 2023; Yan et al., 2024）。在GRADE-Agent中，研究发现为值得关注的方面提供依据，用户补充的信息则为讨论可能的支持方式提供依据。最终总结将已有事实与需要进一步考虑的问题、可能采取的后续措施联系起来。教师仍需核查这些解释，并判断所建议的措施是否适合该学习者。'),
        ('This process creates a distinction between the basis of a suggestion and its educational quality. Relevant research findings do not ensure that a generated reply addresses the case accurately or communicates useful advice. The framework therefore treats contextual relevance, consistency with case facts, appropriateness and actionability of suggestions, and clarity as separate aspects of reply quality. Teacher evaluation examines these aspects without treating favourable ratings as evidence of improved student outcomes.',
         '这一过程也要求区分建议的依据与建议本身的教育质量。即使具有相关研究发现，生成的回复也未必能准确回应案例或清楚传达有用的建议。因此，本框架将情境相关性、与案例事实的一致性、建议的适切性与可操作性，以及表达清晰性视为回复质量的不同方面。教师评价分别考察这些方面，而不把较高评分等同于学生学习结果得到改善的证据。'),
        ('Figure 1 brings these relationships together. The empirical stage establishes the findings and their interpretive scope; the conversational stage relates them to the learner\'s described circumstances. Teacher evaluation then examines the quality of the resulting replies. The framework connects these stages through the interpretation of learning needs while keeping their respective claims distinct.',
         '图1将这些关系整合起来。实证分析阶段确立研究发现及其解释范围；对话阶段将研究发现与学习者所描述的具体情境联系起来；随后，教师评价考察最终回复的质量。该框架以对学习需要的理解连接三个阶段，同时区分各阶段能够支持的研究结论。')
    ])
]
zh22 = [
    '诊断支持的教育依据与形成性评价密切相关。形成性评价通过收集和解释学习信息来确定适当的后续措施，而不只是利用评价给学生分类或概括其表现（Black & Wiliam, 2009）。从这一视角看，当学业成绩数据有助于持续理解学习需要并考虑适当的教育回应时，其价值才能得到体现。',
    '然而，仅凭分数并不总能理解学生的表现。相同的成绩水平可能对应不同的学习经历、学习条件或支持需要。例如，成绩下降可能反映持续存在的困难、情境的暂时变化，或关于学生近期学习情况的信息不完整。解释这类案例，需要关注具体情境，并考虑多种解释同时成立的可能性。教师诊断能力同样涉及将学生的情境和学习行为信息，与能够为后续教学决策提供依据的解释联系起来（Klug et al., 2013）。',
    '这些考虑意味着，诊断支持应当关注情境，并允许修正。诊断过程应当允许纳入补充信息并重新考虑暂定解释，而不是根据最初描述给出固定解释。同时，需要区分已经报告的学生情况与仍属于假设的解释。当现有信息不足时，保留不确定性比强行作出确定结论更为适当。',
    '当分析发现用于支持教育决策时，以人为中心的方法尤其重要。教育人工智能研究强调，在AI输出可能影响对学习者的理解或支持方式时，应重视透明性、利益相关者参与和审慎评价（Bond et al., 2024; Khosravi et al., 2022; Zawacki-Richter et al., 2019）。因此，AI生成的建议应当易于理解且能够接受审阅，教师及其他相关人员仍需负责解释个体情境并决定适当的行动。诊断支持的目的在于辅助专业判断，而不是取代专业判断。'
]
for heading, zhheading, pairs in sections:
    hi = next(i for i,p in enumerate(doc.paragraphs) if p.text == heading)
    current = doc.paragraphs
    stop = next(i for i in range(hi+1,len(current)) if current[i].style.name.startswith('Heading'))
    if heading.startswith('2.2'):
        original22 = [p for p in current[hi+1:stop] if p.text]
        pairs.extend((p.text,zh) for p,zh in zip(original22,zh22))
        continue
    # Preserve image and caption after the final explanatory paragraph.
    old = []
    for p in current[hi+1:stop]:
        if p._p.xpath('.//w:drawing') or p.text.startswith('Figure 1.'):
            continue
        old.append(p)
    template = next(p for p in old if p.text)
    pp = deepcopy(template._p.pPr)
    rpr = deepcopy(template.runs[0]._r.rPr) if template.runs else None
    anchor = current[hi]._p
    for en,zh in pairs:
        node = OxmlElement('w:p')
        anchor.addnext(node)
        p = Paragraph(node, doc._body)
        if pp is not None:
            node.append(deepcopy(pp))
        p.paragraph_format.first_line_indent = Inches(.25)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run=p.add_run(en)
        if rpr is not None:
            run._r.insert(0,deepcopy(rpr))
        anchor=node
    for p in old:
        p._p.getparent().remove(p._p)
doc.save(OUT)
after=doc.paragraphs
newstart=next(i for i,p in enumerate(after) if p.text=='2. Conceptual Framework')
newend=next(i for i,p in enumerate(after) if p.text=='3. Materials and Methods')
assert outside == [p._p.xml for p in after[:newstart]]+[p._p.xml for p in after[newend:]]
assert tables == [t._tbl.xml for t in doc.tables]
assert images == [s._inline.xml for s in doc.inline_shapes]
assert [p._p.xml for p in original22] == [p._p.xml for p in doc.paragraphs if p.text in [q.text for q in original22]]
new_words=sum(len(p.text.split()) for p in after[newstart:newend])

bi=Document()
sec=bi.sections[0]
sec.top_margin=sec.bottom_margin=Inches(.8)
sec.left_margin=sec.right_margin=Inches(.85)
normal=bi.styles['Normal']
normal.font.name='Times New Roman'
normal.font.size=Pt(11)
normal.paragraph_format.space_after=Pt(7)
normal.paragraph_format.line_spacing=1.15
bi.add_heading('Chapter 2 Conceptual Framework\n第二章 概念框架',0)
for heading,zhheading,pairs in sections:
    bi.add_heading(heading+'\n'+zhheading,1)
    for en,zh in pairs:
        p=bi.add_paragraph(en)
        p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent=Inches(.25)
        p=bi.add_paragraph(zh)
        p.paragraph_format.first_line_indent=Pt(22)
        for r in p.runs:
            r._r.get_or_add_rPr().get_or_add_rFonts().set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia','宋体')
for sty in ('Title','Heading 1'):
    bi.styles[sty].font.name='Times New Roman'
    bi.styles[sty].font.color.rgb=__import__('docx').shared.RGBColor(0,0,0)
    bi.styles[sty].font.size=Pt(16 if sty=='Title' else 12)
    bi.styles[sty].paragraph_format.line_spacing=1.4
for element in (bi._element,bi.styles.element):
    for border in element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
for p in bi.paragraphs:
    if p.style.name in ('Title','Heading 1'):
        for run in p.runs:
            run._r.get_or_add_rPr().get_or_add_rFonts().set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia','宋体')
bi.save(BILINGUAL)
data={'old_chapter2_words':old_words,'new_chapter2_words':new_words,'outside_chapter2_unchanged':True,'tables_and_images_unchanged':True,'section22_unchanged':True,'sections':sections}
(QA/'revision.json').write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in data.items() if k!='sections'},ensure_ascii=True))
print(OUT)
print(BILINGUAL)

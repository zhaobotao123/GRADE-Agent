from pathlib import Path
from copy import deepcopy
import json,urllib.request,subprocess,sys,hashlib
from docx import Document
BASE=Path(__file__).resolve().parents[1]
F=Path('D:/Users/84130/Desktop/投稿/转投修改_20261022')
OUT=BASE/'outputs/hu_publication_v36_20261028';OUT.mkdir(parents=True,exist_ok=True)
doi='10.1016/j.stueduc.2026.101701'
try:
 req=urllib.request.Request('https://api.crossref.org/works/'+doi,headers={'User-Agent':'GRADEAgentReferenceCheck/1.1'})
 with urllib.request.urlopen(req,timeout=18)as r:metadata=json.load(r)['message']
 (OUT/'crossref_hu.json').write_text(json.dumps(metadata,ensure_ascii=False,indent=2),encoding='utf-8')
 print(json.dumps({k:metadata.get(k)for k in ['title','author','volume','page','article-number','published','published-online','published-print']},ensure_ascii=False),flush=True)
 assert 'mathematics learning efficiency' in metadata['title'][0].lower()
except Exception as exc:
 metadata={'source':'User-provided ScienceDirect publication screenshot','year':2026,'volume':'91','article-number':'101701','DOI':doi}
 (OUT/'screenshot_metadata.json').write_text(json.dumps(metadata,ensure_ascii=False,indent=2),encoding='utf-8')
 print('CrossRef unavailable; using clearly legible publisher screenshot:',type(exc).__name__,flush=True)
d=Document(F/'GRADE-Agent_SEE_Checked_v35.docx')
before=[p._p.xml for p in d.paragraphs]
changes=[]
for i,p in enumerate(d.paragraphs):
 if 'Hu et al., in press' in p.text:
  old=p.text
  for r in p.runs:r.text=r.text.replace('Hu et al., in press','Hu et al., 2026')
  assert p.text==old.replace('Hu et al., in press','Hu et al., 2026')
  changes.append(i)
 if p.text.startswith('Hu, W., Zhang, Y., Zhang, Y., & Yang, X. (in press)'):
  prop=deepcopy(p.runs[0]._r.rPr)
  p.clear()
  for t,it in [('Hu, W., Zhang, Y., Zhang, Y., & Yang, X. (2026). Evaluating and diagnosing mathematics learning efficiency with an intelligent agent: A double machine learning framework. ',False),('Studies in Educational Evaluation, 91',True),(', Article 101701. https://doi.org/'+doi,False)]:
   r=p.add_run(t)
   if prop is not None:r._r.insert(0,deepcopy(prop))
   r.italic=it
  changes.append(i)
assert len(changes)==2,changes
assert all(old==p._p.xml for i,(old,p)in enumerate(zip(before,d.paragraphs))if i not in changes)
assert not any('in press'in p.text for p in d.paragraphs)
target=F/'GRADE-Agent_SEE_Checked_v36.docx';d.save(target)
note='''# Hu文献出版信息更新 v36

根据作者提供的ScienceDirect正式页面，Hu、Zhang、Zhang与Yang的论文已线上发表，有完整卷号、文章号及DOI。

Hu, W., Zhang, Y., Zhang, Y., & Yang, X. (2026). Evaluating and diagnosing mathematics learning efficiency with an intelligent agent: A double machine learning framework. Studies in Educational Evaluation, 91, Article 101701. https://doi.org/10.1016/j.stueduc.2026.101701

正文Hu et al., in press改为Hu et al., 2026；文后补齐2026、91、101701和DOI。卷期标示为December 2026与提前线上发布并不矛盾。
仅两处引用段落修改；其他正文、实验数值、图表、公式及补充材料不变。v35核查记录中Hu的“在印状态待确认”现由本记录更新。截图可确认书目信息，不据此声称重新核查了论文全文。
'''
(F/'Hu文献出版信息更新_v36.md').write_text(note,encoding='utf-8')
print('Saved',target,flush=True)
renderer=Path('C:/Users/84130/.codex/plugins/cache/openai-primary-runtime/documents/26.1025.11240/skills/documents/render_docx.py')
subprocess.run([sys.executable,str(renderer),str(target),'--output_dir',str(OUT/'render'),'--emit_pdf'],check=True)

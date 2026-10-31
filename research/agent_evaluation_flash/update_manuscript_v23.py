"""Insert the author's approved Section 5.5 verbatim, using v21 as the base."""
from pathlib import Path
from copy import deepcopy
import json
from docx import Document
from docx.shared import Pt,Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

base=Path('D:/Users/84130/Desktop/投稿/转投修改_20261022')
source=base/'GRADE-Agent_SEE_Teacher_Evaluation_v21.docx'
target=base/'GRADE-Agent_SEE_Teacher_Evaluation_v23.docx'
d=Document(source);ps=list(d.paragraphs)
start=next(p for p in ps if p.text.startswith('5.5 Teacher-rated'))
end=next(p for p in ps if p.text=='6. Discussion')
before=[p._p.xml for p in ps[:ps.index(start)]]
after=[p._p.xml for p in ps[ps.index(end):]]
equations=[p._p.xml for p in ps if p._p.xpath('.//m:oMath')]
figures=[s._inline.xml for s in d.inline_shapes]
table=deepcopy(d.tables[1]._tbl)
bodypr=deepcopy(ps[146]._p.pPr)
captionpr=deepcopy(ps[147]._p.pPr)
figcaption=next(p for p in ps if p.text.startswith('Figure 5. Teacher-rated'))
figimage=figcaption._p.getprevious()
node=start._p.getnext()
while node is not figimage:
 nextnode=node.getnext();node.getparent().remove(node);node=nextnode
heading='5.5. Teacher Evaluation of Educational Relevance and Learning-Support Recommendations'
start.clear();start.add_run(heading)

paragraphs=[
"The value of an educational support system lies not simply in identifying factors associated with student achievement, but also in helping educators interpret students' circumstances and consider appropriate ways to support their learning. The teacher evaluation therefore examined whether GRADE-Agent could provide responses that were relevant to the learning situation described, consistent with the case information, clear in their explanations, and useful for deciding what support might be offered.",
"The clearest difference between GRADE-Agent and the ordinary multi-turn LLM concerned the appropriateness and actionability of learning-support recommendations. GRADE-Agent received a mean rating of 4.75 out of 5, compared with 3.61 for the ordinary LLM. The mean paired difference across the twelve cases was 1.14 points (95% bootstrap CI [0.86, 1.45]), with higher ratings for GRADE-Agent in every case. This result suggests that teachers found GRADE-Agent's recommendations more suitable for the educational situations presented and more useful in identifying possible next steps for learning support.",
"The results also point to differences in how the systems addressed the broader educational context. GRADE-Agent received higher ratings for contextual relevance (4.03 versus 3.53), indicating that its responses were judged to fit the circumstances described in the cases more closely. Its ratings were also higher for consistency with case facts (4.10 versus 3.91) and clarity (4.14 versus 3.79). Although the difference was smaller for factual consistency, the overall pattern suggests that teachers valued responses that connected the information provided about a student with relevant educational considerations and potential support actions, rather than offering general advice alone.",
"The two teachers' ratings showed a similar pattern for the primary outcome. Teacher 1 gave GRADE-Agent a mean rating of 4.90 for suggestion appropriateness and actionability, compared with 3.57 for the ordinary LLM; the corresponding ratings from Teacher 2 were 4.60 and 3.65. Agreement between the teachers was also stronger for contextual relevance and suggestion appropriateness than for factual consistency and clarity. Quadratic-weighted κ values were 0.794, 0.797, 0.178, and 0.116, respectively. This difference suggests that the educational relevance and practical usefulness of recommendations were more consistently judged than some aspects of factual consistency and expression, highlighting the importance of clear and operational evaluation criteria.",
"Taken together, the teacher ratings suggest that GRADE-Agent's main contribution lies in connecting the analysis of achievement-related factors with educationally relevant interpretations and practical learning-support recommendations. Rather than treating the identification of relevant factors as the endpoint, the system provides a way to present those findings in relation to a learner's described circumstances and to consider possible next steps. This offers a promising direction for making empirical research on student achievement more accessible to educational discussion and support planning.",
]
note="Note. Ratings use a five-point scale and are averaged across two teachers and 60 replies per system (12 cases × 5 runs). Paired differences are calculated from case-level mean ratings; positive values favour GRADE-Agent. Confidence intervals are percentile bootstrap intervals based on 50,000 resamples of the 12 paired cases. Appropriateness and actionability of suggestions was the prespecified primary outcome. Weighted κ denotes quadratic-weighted Cohen's kappa based on the two teachers' ratings."
caption='Table 2. Teacher ratings of the educational quality of final replies'

def insert(s,cap=False):
 p=d.add_paragraph()
 figimage.addprevious(p._p)
 if p._p.pPr is not None:p._p.remove(p._p.pPr)
 props=captionpr if cap else bodypr
 if props is not None:p._p.insert(0,deepcopy(props))
 p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
 p.paragraph_format.keep_with_next=False
 p.paragraph_format.keep_together=False
 r=p.add_run(s);r.font.name='Times New Roman';r.font.size=Pt(10 if cap else 12)
 return p

insert(paragraphs[0]);insert(paragraphs[1])
p=insert(caption,True);p.runs[0].bold=True;p.paragraph_format.keep_with_next=True
figimage.addprevious(table)
from docx.table import Table
t=Table(table,d._body)
rows=[
['Evaluation dimension','Ordinary LLM','GRADE-Agent','Paired difference [95% CI]','Weighted κ'],
['Contextual relevance','3.53','4.03','0.50 [0.35, 0.63]','0.794'],
['Consistency with case facts','3.91','4.10','0.19 [0.10, 0.28]','0.178'],
['Appropriateness and actionability of suggestions','3.61','4.75','1.14 [0.86, 1.45]','0.797'],
['Clarity','3.79','4.14','0.35 [0.24, 0.48]','0.116'],
]
widths=[2.15,.75,.80,1.65,.85]
for c,w in zip(t.columns,widths):c.width=Inches(w)
for ri,(row,vals) in enumerate(zip(t.rows,rows)):
 for ci,(cell,value,w) in enumerate(zip(row.cells,vals,widths)):
  cell.width=Inches(w)
  p=cell.paragraphs[0];p.clear();p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 else WD_ALIGN_PARAGRAPH.CENTER
  r=p.add_run(value);r.font.name='Times New Roman';r.font.size=Pt(10);r.bold=(ri==0)
p=insert(note,True)
p.clear();r=p.add_run('Note.');r.italic=True;r.font.size=Pt(10);r.font.name='Times New Roman'
r=p.add_run(note[len('Note.'):]);r.font.size=Pt(10);r.font.name='Times New Roman'
for s in paragraphs[2:]:insert(s)

allp=list(d.paragraphs);newstart=next(p for p in allp if p.text==heading);newend=next(p for p in allp if p.text=='6. Discussion')
assert before==[p._p.xml for p in allp[:allp.index(newstart)]]
assert after==[p._p.xml for p in allp[allp.index(newend):]]
assert figures==[s._inline.xml for s in d.inline_shapes]
assert equations==[p._p.xml for p in allp if p._p.xpath('.//m:oMath')]
newsection=[p.text for p in allp[allp.index(newstart)+1:allp.index(newend)] if p.text and not p.text.startswith('Figure 5.')]
assert newsection==[paragraphs[0],paragraphs[1],caption,note,*paragraphs[2:]]
# Avoid a nearly empty page containing only the short AI declaration after reflow.
# This changes pagination only, not any wording outside the approved section.
next(p for p in d.paragraphs if p.text=='References').paragraph_format.page_break_before=False
d.save(target)
qa=Path(__file__).resolve().parents[1]/'outputs/manuscript_section55_v23_20261027/qa'
qa.mkdir(parents=True,exist_ok=True)
(qa/'verification.json').write_text(json.dumps({'base':str(source),'output':str(target),'author_text_verbatim':True,'outside_section_55_text_unchanged':True,'pagination_adjustment':'Removed forced References page break to avoid a nearly empty AI-declaration page.','all_figures_equations_unchanged':True,'table_data':rows},ensure_ascii=False,indent=2),encoding='utf-8')
print(target)

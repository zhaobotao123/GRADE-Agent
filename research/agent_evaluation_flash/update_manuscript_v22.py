from pathlib import Path
from docx import Document
from docx.shared import Inches
import json,re

folder=Path('D:/Users/84130/Desktop/投稿/转投修改_20261022')
source=folder/'GRADE-Agent_SEE_Teacher_Evaluation_v21.docx'
target=folder/'GRADE-Agent_SEE_Teacher_Evaluation_v22.docx'
d=Document(source)
equations=[p._p.xml for p in d.paragraphs if p._p.xpath('.//m:oMath')]
media=[s._inline.xml for s in d.inline_shapes]
removed=[];edited=[]
for p in list(d.paragraphs):
 s=p.text;t=s
 if s.startswith('Inter-rater agreement differed substantially') or s.startswith('The comparison also separates score differences'):
  removed.append(s);p._p.getparent().remove(p._p);continue
 for sentence in [
  ' Agreement between teachers varied across dimensions.',
  " Quadratic-weighted Cohen's kappa and exact agreement were calculated separately for each criterion across the 120 matched replies.",
  ' Mean system differences and inter-rater agreement answer different questions and are reported separately.',
  " Weighted κ is quadratic-weighted Cohen's kappa across the 120 replies.",
  ' Agreement was limited for case-fact consistency and clarity, despite higher mean scores for GRADE-Agent.',
  ' Agreement between teachers varied across criteria.',
 ]:t=t.replace(sentence,'')
 t=t.replace('Results are reported as mean ratings, paired differences with case-cluster bootstrap intervals, and criterion-specific agreement between teachers.','Results are reported as mean ratings and paired differences with case-cluster bootstrap intervals.')
 if t!=s:
  edited.append({'before':s,'after':t})
  # These paragraphs contain only ordinary text, not equations or drawings.
  assert not p._p.xpath('.//m:oMath | .//w:drawing')
  props=p.runs[0]._r.rPr if p.runs else None
  from copy import deepcopy
  props=deepcopy(props) if props is not None else None
  p.clear();r=p.add_run(t)
  if props is not None:r._r.insert(0,props)

table=d.tables[1]
assert 'κ' in table.cell(0,4).text
for row in table.rows:
 cell=row._tr.tc_lst[-1];row._tr.remove(cell)
table._tbl.tblGrid.remove(table._tbl.tblGrid.gridCol_lst[-1])
widths=[2.55,.90,.90,1.85]
for col,w in zip(table.columns,widths):col.width=Inches(w)
for row in table.rows:
 for cell,w in zip(row.cells,widths):cell.width=Inches(w)
assert equations==[p._p.xml for p in d.paragraphs if p._p.xpath('.//m:oMath')]
assert media==[s._inline.xml for s in d.inline_shapes]
text='\n'.join(p.text for p in d.paragraphs)+'\n'+'\n'.join(c.text for t in d.tables for r in t.rows for c in r.cells)
assert not re.search(r'inter-rater|kappa|κ|agreement between teachers|weighted agreement',text,re.I)
d.save(target)
qa=Path(__file__).resolve().parents[1]/'outputs/manuscript_teacher_evaluation_v22_20261027/qa'
qa.mkdir(parents=True,exist_ok=True)
(qa/'changes.json').write_text(json.dumps({'source':str(source),'output':str(target),'removed_paragraphs':removed,'edited_paragraphs':edited,'table_last_column_removed':True,'images_equations_unchanged':True,'original_ratings_and_analysis_unchanged':True},ensure_ascii=False,indent=2),encoding='utf-8')
print(target)

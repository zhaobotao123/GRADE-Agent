from pathlib import Path
from docx import Document
ROOT=Path(__file__).resolve().parent
SOURCE=Path('D:/Users/84130/Desktop/投稿/转投修改_20261022/GRADE-Agent_SEE_Educational_Contribution_v38.docx')
d=Document(SOURCE)
lines=[f'{i}\t{p.style.name}\t{p.text}' for i,p in enumerate(d.paragraphs)]
(ROOT/'v38_paragraphs.txt').write_text('\n'.join(lines),encoding='utf-8')
for i,p in enumerate(d.paragraphs):
    if (45<i<110 or 140<i<240): print(i,p.style.name,p.text)
